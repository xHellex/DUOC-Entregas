# -*- coding: utf-8 -*-
"""
Genera el informe de la Experiencia 3 - Semana 7 (tolerancia a fallos y
arquitectura de eventos con Kafka, actividad FORMATIVA grupal) sobre la
plantilla oficial DUOC (PBY2202_EFT_Plantilla_Informe_PDF.docx), reutilizando
sus estilos y dejando portada, indice/TOC y aviso legal intactos. Mismo
formato que el informe de la Semana 6.

Uso:  python scripts/generar_informe_duoc.py
Requiere la plantilla en el nivel superior de la carpeta S7 (ultraDocumentBody):
  ../PBY2202_EFT_Plantilla_Informe_PDF.docx
Salida: Informe_Exp3_S7_Grupo3.docx (en la raiz del proyecto).
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EVID = os.path.join(RAIZ, "evidencia")
PLANTILLA = os.path.join(os.path.dirname(RAIZ), "PBY2202_EFT_Plantilla_Informe_PDF.docx")
SALIDA = os.path.join(RAIZ, "Informe_Exp3_S7_Grupo3.docx")

INTEGRANTES = "Felipe Peñaloza"   # Grupo 3 (integrante unico)

if not os.path.exists(PLANTILLA):
    raise SystemExit(f"No se encontro la plantilla en: {PLANTILLA}")

doc = Document(PLANTILLA)
body = doc.element.body

anchor = None
for p in doc.paragraphs:
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
        anchor = p._p
        break
if anchor is None:
    raise SystemExit("No se encontro el parrafo con el salto de seccion (sectPr).")

def set_text(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for r in paragraph.runs[1:]:
        r.text = ""

paras = doc.paragraphs
set_text(paras[16], "Desarrollo Backend III (PBY2203)")
set_text(paras[17], "Configurando tolerancia a fallos y arquitectura de eventos con microservicios en la nube")

p18 = paras[18]
p18.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p18.add_run("Arquitectura de eventos con Kafka para el Banco XYZ")
r.font.size = Pt(14)

p19 = paras[19]
p19.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p19.add_run("Experiencia 3 - Semana 7 (actividad formativa grupal)\n"
                f"Grupo 3 - {INTEGRANTES}\n"
                "Facilitador disciplinar: Ana Karina Villagran Ibarra")
r.font.size = Pt(11)

INICIO_EJEMPLO, FIN_EJEMPLO = 26, 48
elems_a_borrar = []
idx = 0
for child in list(body):
    tag = child.tag.split("}")[-1]
    if tag in ("p", "tbl"):
        if INICIO_EJEMPLO <= idx <= FIN_EJEMPLO:
            elems_a_borrar.append(child)
        idx += 1
for el in elems_a_borrar:
    el.getparent().remove(el)

def _mover(el):
    anchor.addprevious(el)

def h1(t):
    p = doc.add_paragraph(t, style="Heading 1"); _mover(p._p); return p

def h2(t):
    p = doc.add_paragraph(t, style="Heading 2"); _mover(p._p); return p

def par(t, bold=False):
    p = doc.add_paragraph(); rr = p.add_run(t); rr.bold = bold; _mover(p._p); return p

def bullet(t):
    p = doc.add_paragraph(style="List Paragraph"); p.add_run("•  " + t); _mover(p._p); return p

def code(t):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    rr = p.add_run(t)
    rr.font.name = "Consolas"; rr.font.size = Pt(8.5)
    rr.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    _mover(p._p); return p

def salto_pagina():
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); _mover(p._p)

def imagen(nombre, ancho=6.0, pie=None):
    ruta = os.path.join(EVID, nombre)
    if not os.path.exists(ruta):
        par(f"[FALTA IMAGEN: {nombre}]", bold=True); return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=Inches(ancho)); _mover(p._p)
    if pie:
        cp = doc.add_paragraph(style="titulo-tabla"); cp.add_run(pie); _mover(cp._p)

_tn = [0]
def tabla(titulo, cabeceras, filas):
    _tn[0] += 1
    tp = doc.add_paragraph(style="titulo-tabla")
    tp.add_run(f"Tabla {_tn[0]}: {titulo}"); _mover(tp._p)
    t = doc.add_table(rows=1, cols=len(cabeceras)); t.style = "Table Grid"
    for i, c in enumerate(cabeceras):
        cell = t.rows[0].cells[i]; cell.text = c
        for p in cell.paragraphs:
            for rr in p.runs:
                rr.bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = str(v)
    _mover(t._tbl)
    sp = doc.add_paragraph(); _mover(sp._p)

# ================================================================
#  CONTENIDO
# ================================================================
salto_pagina()

h1("1. Introduccion")
par("Este informe corresponde a la Experiencia 3, Semana 7. Continua el ecosistema de "
    "microservicios del Banco XYZ construido en la Semana 6 (Config Server, Eureka, API "
    "Gateway y los microservicios de cuentas, transacciones y tarjetas) y le agrega dos "
    "capacidades: tolerancia a fallos con Resilience4j y una arquitectura de eventos "
    "asincrona con Apache Kafka, aplicada a las transacciones del sistema.")
par("El informe esta organizado en el mismo orden que los cuatro criterios de la pauta "
    "formativa. Cada afirmacion se respalda con salidas reales de ejecucion guardadas en la "
    "carpeta evidencia/: se levanto un broker Kafka en Docker, los seis modulos Spring Boot, "
    "y se probaron el flujo normal, la caida del broker, la caida de un consumidor, la "
    "escalabilidad con dos instancias de un mismo grupo y la apertura real de un Circuit "
    "Breaker.")
par("Ademas, la retroalimentacion de la Semana 6 pidio explicar con mayor detalle el rol del "
    "API Gateway, por tratarse de un componente adicional a lo solicitado. Esa explicacion "
    "esta en la seccion 'Rol del API Gateway' del Criterio 1.")

h1("2. Resultado de aprendizaje abordado")
par("RA3. Implementa arquitecturas de microservicios resilientes y seguras en la nube, "
    "aplicando patrones de tolerancia a fallos y arquitecturas de eventos, utilizando el "
    "ecosistema Spring Cloud y sistemas de mensajeria asincrona.")

h1("3. Caso de uso y patron elegido")
par("Cuando se registra una transaccion, otros servicios necesitan reaccionar: ms-cuentas "
    "debe actualizar el saldo y ms-tarjetas debe registrar el movimiento. La pregunta de "
    "diseno es como se entera cada uno.")
par("Patron elegido: Publish/Subscribe (publicador/suscriptor) sobre Apache Kafka. "
    "ms-transacciones publica un evento TransaccionCreadaEvent en el topico transacciones; "
    "cada servicio interesado se suscribe con su propio grupo de consumidores. Una "
    "transaccion es un hecho del negocio que interesa a varios servicios a la vez, y el "
    "productor no necesita saber quien la consume.")
tabla("Alternativas de diseno evaluadas",
      ["Alternativa", "Como funcionaria", "Por que se descarto / se eligio"],
      [["Llamadas REST sincronas entre microservicios",
        "ms-transacciones llama a ms-cuentas y ms-tarjetas",
        "Descartada: acoplamiento temporal; si un servicio esta caido la transaccion falla, "
        "y agregar un consumidor obliga a modificar el productor"],
       ["Cola punto a punto (JMS Queue)",
        "Un mensaje lo recibe un solo consumidor",
        "Descartada: solo uno de los dos servicios recibiria cada transaccion"],
       ["Event Sourcing / CQRS",
        "El estado se reconstruye desde el log de eventos",
        "Descartada: complejidad desproporcionada para el alcance de la actividad"],
       ["Publish/Subscribe (Kafka)", "Un evento, N suscriptores independientes",
        "Elegida: desacopla productor y consumidores, cada uno recibe una copia completa y "
        "se puede escalar"]])
par("Kafka o JMS: la pauta acepta ambos. Se eligio Kafka porque persiste los mensajes en un "
    "log con offsets por grupo (permite reintentar y ponerse al dia tras una caida), escala "
    "por particiones y deja evidencia visual con Kafka-UI. El costo es requerir un broker "
    "externo, resuelto con un docker-compose.yml. JMS con ActiveMQ embebido habria sido mas "
    "simple de arrancar, pero no permite demostrar la escalabilidad horizontal por "
    "particiones.")

h1("4. Criterio 1: Define la arquitectura de eventos, alineada con los patrones de diseno y adecuada al caso de uso")
par("Redaccion de la pauta: 'Define la arquitectura de eventos a utilizar, alineada con los "
    "patrones de diseno seleccionados, y es adecuada para el caso de uso'.")
tabla("Decisiones de diseno y su justificacion",
      ["Decision", "Justificacion"],
      [["Un unico topico transacciones, con 3 particiones",
        "La particion es la unidad de paralelismo: un grupo puede tener hasta 3 instancias "
        "activas"],
       ["Key del mensaje = cuentaId",
        "Las transacciones de una misma cuenta caen siempre en la misma particion y "
        "conservan su orden"],
       ["Un groupId distinto por microservicio (grupo-cuentas, grupo-tarjetas)",
        "Cada grupo recibe TODOS los eventos (broadcast, esencia de Pub/Sub); dentro de un "
        "grupo las instancias se reparten las particiones"],
       ["El contrato del evento es JSON puro, sin el header __TypeId__",
        "Los servicios no comparten codigo: cada uno tiene su propia copia del record del "
        "evento"],
       ["Cada servicio define su propia copia del evento",
        "Mantiene el desacople; no hay una libreria compartida que acople las versiones"]])
h2("4.1 Rol del API Gateway")
par("El API Gateway (puerto 8080) NO forma parte del flujo de eventos: es un componente "
    "adicional respecto de lo solicitado, heredado de la Semana 6. Su funcion es ser la "
    "puerta de entrada unica del sistema para los clientes HTTP:")
bullet("Los clientes conocen una sola direccion (localhost:8080) y no los puertos de cada "
       "microservicio.")
bullet("Enruta por nombre logico (lb://ms-transacciones), resolviendo la ubicacion real de "
       "cada servicio a traves de Eureka; si un servicio cambia de puerto o se levanta otra "
       "instancia, los clientes no se enteran.")
bullet("Delega la autenticacion en los microservicios: reenvia el header Authorization: "
       "Bearer <JWT> y cada servicio valida el token por su cuenta.")
par("En esta semana el Gateway es el punto por el que ingresa el POST /transacciones que "
    "origina el evento, pero la comunicacion entre microservicios (productor -> Kafka -> "
    "consumidores) ocurre sin pasar por el Gateway: es asincrona y va directamente contra el "
    "broker. Por eso el diagrama del Criterio 2 lo muestra a la izquierda, antes del "
    "productor, y no en el centro del flujo.")

h1("5. Criterio 2: Diagrama representativo de la arquitectura, con topicos, mensajes y eventos")
par("Redaccion de la pauta: 'Elabora un diagrama representativo de la arquitectura elegida "
    "de forma completa con los topicos/mensajes/eventos de la solucion, con una estructura "
    "visual organizada'.")
imagen("diagrama_arquitectura_eventos.png", ancho=6.5,
       pie="Arquitectura de eventos Publish/Subscribe con Kafka: cliente -> gateway -> "
           "productor -> topico con 3 particiones -> grupos de consumidores")
par("El diagrama se lee de izquierda a derecha en tres bloques: entrada y productor "
    "(cliente, gateway, ms-transacciones), broker Kafka (topico con sus particiones y el "
    "mensaje) y consumidores (un recuadro por grupo). La franja inferior muestra los "
    "componentes de soporte.")
tabla("Topicos, mensajes y eventos de la solucion",
      ["Elemento", "Detalle"],
      [["Topico", "transacciones - 3 particiones, factor de replicacion 1"],
       ["Evento", "TransaccionCreadaEvent - hecho 'se creo una transaccion'"],
       ["Campos del mensaje",
        "transaccionId (Long), cuentaId (Long), monto (Double), tipo (credito / debito), "
        "fecha (texto ISO, ej. 2026-09-24)"],
       ["Key", "cuentaId"],
       ["Productor", "ms-transacciones (TransaccionEventProducer)"],
       ["Consumidor 1", "ms-cuentas, groupId = grupo-cuentas: credito suma, debito resta al "
                         "saldo"],
       ["Consumidor 2", "ms-tarjetas, groupId = grupo-tarjetas: registra el movimiento "
                         "(instancias :8083 y :8093)"]])
par("Mensaje real leido directamente desde el topico (extracto de evidencia/01_kafka_topico.txt):")
code('Offset:0  NO_HEADERS  101  {"transaccionId":4,"cuentaId":101,"monto":500.0,"tipo":"credito","fecha":"2026-09-24"}\n'
     'Offset:1  NO_HEADERS  101  {"transaccionId":5,"cuentaId":101,"monto":300.0,"tipo":"debito","fecha":"2026-09-24"}')

h1("6. Criterio 3: Tolerancia a fallos con Resilience4j demostrando resiliencia ante fallos")
par("La tolerancia a fallos se aplica en dos capas: el Circuit Breaker de Resilience4j en "
    "cada microservicio, y las propiedades de resiliencia propias de la arquitectura de "
    "eventos.")
h2("6.1 Capa 1: Circuit Breaker con Resilience4j")
tabla("Circuit Breaker por microservicio",
      ["Microservicio", "Instancia", "Configuracion"],
      [["ms-cuentas", "cuentasCB", "ventana 10 llamadas, umbral de fallas 50%, 10s abierto"],
       ["ms-transacciones", "transaccionesCB",
        "ventana 10 llamadas, umbral de fallas 50%, 10s abierto"],
       ["ms-tarjetas", "tarjetasCB", "ventana 10 llamadas, umbral de fallas 50%, 10s abierto"]])
code('@GetMapping\n'
     '@CircuitBreaker(name = "transaccionesCB", fallbackMethod = "listarFallback")\n'
     'public List<Transaccion> listar() { return service.listar(); }\n\n'
     'public List<Transaccion> listarFallback(Throwable t) {\n'
     '    log.warn("[FALLBACK] listar transacciones: {}", t.getMessage());\n'
     '    return List.of();\n'
     '}')
par("Demostracion con una falla real: se inyecto temporalmente una excepcion en "
    "TransaccionService.listar() (revertida despues) y se hicieron 10 consultas. Extracto de "
    "evidencia/05_circuit_breaker_con_eventos.txt:")
tabla("Ciclo completo del Circuit Breaker con falla real",
      ["Paso", "Resultado observado"],
      [["10 consultas con el servicio fallando",
        "El cliente recibe 200 [] (fallback) en las 10; nunca un error 500"],
       ["Estado tras las 10 fallas", "failureRate 100%, failedCalls 10, state OPEN"],
       ["Consulta numero 11 con el circuito abierto",
        "notPermittedCalls 1; el log indica 'CircuitBreaker transaccionesCB is OPEN and does "
        "not permit further calls'"],
       ["Con el circuito abierto, POST /transacciones",
        "201; el evento se publica y ms-cuentas actualiza el saldo de la cuenta 101 a 5080.0"],
       ["Se revierte la falla y se reinicia",
        "state CLOSED, GET /transacciones vuelve a devolver los datos reales"]])
par("La cuarta fila muestra el aislamiento de fallas: que el circuito de las consultas este "
    "abierto no afecta al registro de transacciones ni a su evento.")
h2("6.2 Capa 2: resiliencia propia de la arquitectura de eventos")
par("Extracto de evidencia/03_resiliencia_consumidor_y_broker.txt:")
tabla("Escenarios de falla de la arquitectura de eventos",
      ["Escenario", "Que se hizo", "Resultado"],
      [["Broker caido", "docker stop del broker y un POST /transacciones",
        "201 en 0,011 s (el send es asincrono). El evento quedo en el buffer del productor y, "
        "al reiniciar el broker, llego a ambos consumidores (saldo de la cuenta 102: 8000 -> "
        "8100)"],
       ["Consumidor caido",
        "Se detuvo ms-tarjetas y se registraron 2 transacciones",
        "Ambas respondieron 201. ms-cuentas (vivo) las proceso (saldo cuenta 103: 12800). El "
        "grupo grupo-tarjetas acumulo lag = 2 (offset 3 de 5)"],
       ["Consumidor se recupera", "Se reinicio ms-tarjetas",
        "Proceso los 2 eventos pendientes al arrancar y el lag volvio a 0 (offset 5 de 5)"]])
par("Un consumidor caido no afecta ni al productor ni a los demas consumidores, y los "
    "eventos esperan en el topico hasta ser procesados.")
h2("6.3 Limitaciones conocidas")
bullet("Si el broker permanece caido mas que delivery.timeout.ms (120s por defecto) o el "
       "productor se cae con eventos en su buffer, esos eventos se pierden: no hay patron "
       "Transactional Outbox.")
bullet("La entrega es at-least-once: si un evento se reentregara, ms-cuentas volveria a "
       "aplicarlo (los consumidores no son idempotentes).")
bullet("Un mensaje que no se pueda deserializar bloquea su particion hasta corregirlo: falta "
       "un ErrorHandlingDeserializer y un topico de mensajes fallidos (dead-letter).")
par("Son mejoras razonables para una version productiva y estan fuera del alcance de la "
    "actividad.")

h1("7. Criterio 4: Integracion de mensajeria asincrona (Kafka), eventos procesados y escalabilidad demostrada")
h2("7.1 Integracion")
par("Se agrego spring-kafka a los tres microservicios. El productor serializa a JSON y no "
    "envia el header de tipo; los consumidores ignoran cualquier header de tipo y leen el "
    "mensaje en su propia clase:")
code("# config-repo/ms-transacciones.yml  (productor)\n"
     "spring.kafka:\n"
     "  bootstrap-servers: localhost:9092\n"
     "  producer:\n"
     "    key-serializer: org.apache.kafka.common.serialization.StringSerializer\n"
     "    value-serializer: org.springframework.kafka.support.serializer.JsonSerializer\n"
     "    properties:\n"
     "      spring.json.add.type.headers: false")
code("// Productor: publica con key = cuentaId\n"
     'kafkaTemplate.send("transacciones", String.valueOf(evento.cuentaId()), evento);\n\n'
     "// Consumidor (ms-cuentas): un grupo propio\n"
     '@KafkaListener(topics = "transacciones", groupId = "grupo-cuentas")\n'
     "public void consumir(TransaccionCreadaEvent evento) { /* suma o resta al saldo */ }")
h2("7.2 Eventos correctamente procesados")
par("Flujo real via API Gateway sobre la cuenta 101 (extracto de "
    "evidencia/02_flujo_eventos_kafka.txt):")
code("Saldo ANTES                       5000.0\n"
     "POST /transacciones credito 500   -> 201\n"
     "Saldo DESPUES                     5500.0    (ms-cuentas consumio el evento)\n"
     "POST /transacciones debito 300    -> 201\n"
     "Saldo DESPUES                     5200.0    (5000 + 500 - 300)\n\n"
     "[KAFKA-PRODUCER] Publicando evento en topico 'transacciones': TransaccionCreadaEvent"
     "[transaccionId=4, cuentaId=101, monto=500.0, tipo=credito, ...]\n"
     "[KAFKA-CONSUMER cuentas]  Evento recibido: TransaccionCreadaEvent[transaccionId=4, ...]"
     "   20:38:06.943\n"
     "[KAFKA-CONSUMER cuentas]  Saldo actualizado para cuenta 101: nuevo saldo 5500.0\n"
     "[KAFKA-CONSUMER tarjetas] Evento recibido: TransaccionCreadaEvent[transaccionId=4, ...]"
     "   20:38:06.943\n"
     "[KAFKA-CONSUMER tarjetas] Registrando movimiento de credito por 500.0 en cuenta 101")
par("Ambos consumidores recibieron el mismo evento en el mismo milisegundo (20:38:06.943): "
    "es el broadcast de Publish/Subscribe, consecuencia de que cada uno use su propio "
    "groupId.")
h2("7.3 Escalabilidad demostrada")
par("Se levanto una segunda instancia de ms-tarjetas (puerto 8093, mismo groupId) y se "
    "publicaron 12 eventos de 5 cuentas distintas. Kafka rebalanceo las particiones entre "
    "las dos instancias (extracto de evidencia/04_escalabilidad_particiones.txt):")
tabla("Reparto de 12 eventos entre las instancias de un mismo grupo",
      ["Instancia", "Particiones asignadas", "Eventos recibidos", "Cuentas atendidas"],
      [["A (ms-tarjetas :8083)", "2", "7", "101, 103, 105"],
       ["B (ms-tarjetas :8093)", "0 y 1", "5", "102, 107"],
       ["Total A + B", "-", "12 (0 duplicados)", "-"],
       ["ms-cuentas (otro grupo)", "0, 1 y 2", "12 (todos)", "-"]])
par("Cada evento lo proceso una sola instancia del grupo, sin duplicados ni perdidas, y una "
    "misma cuenta siempre cayo en la misma instancia (por la key cuentaId). El otro grupo "
    "recibio los 12: escalar un grupo reparte el trabajo; agregar un grupo nuevo agrega un "
    "suscriptor sin tocar al productor. La particion 1 no recibio eventos con estas cinco "
    "cuentas (la distribucion depende del hash de la key).")
h2("7.4 Pruebas automatizadas")
par("Se ejecuto mvn clean test en los seis modulos: 12 pruebas, 0 fallas (extracto en "
    "evidencia/00_compilacion_tests.txt). Cada modulo incluye una prueba de carga de "
    "contexto Spring; ademas, ms-transacciones tiene una prueba de contrato del mensaje "
    "publicado (JSON puro, fecha ISO, sin header __TypeId__) y ms-cuentas y ms-tarjetas una "
    "del mensaje consumido (lectura en su propia clase, ignorando un __TypeId__ ajeno). Se "
    "comprobo que la prueba de consumo FALLA con la configuracion defectuosa original y "
    "reproduce el mismo error de produccion.")
comp = os.path.join(EVID, "00_compilacion_tests.txt")
if os.path.exists(comp):
    with open(comp, encoding="utf-8", errors="replace") as f:
        code(f.read().strip())
h2("7.5 Evidencia visual del broker (Kafka-UI)")
imagen("06_kafkaui_topico_particiones.png", ancho=6.5,
       pie="Kafka-UI: pestana Overview del topico transacciones, con Partitions: 3 y el "
           "detalle de offsets por particion")
imagen("07_kafkaui_mensajes.png", ancho=6.5,
       pie="Kafka-UI: pestana Messages del topico transacciones, con los eventos JSON reales "
           "(transaccionId, cuentaId, monto, tipo, fecha)")
imagen("08_kafkaui_consumers.png", ancho=6.5,
       pie="Kafka-UI: pantalla Consumers, con grupo-cuentas y grupo-tarjetas en estado STABLE")

h1("8. Problemas encontrados durante la verificacion")
par("Al ejecutar el sistema completo (no solo revisar el codigo) aparecieron cinco defectos, "
    "todos corregidos:")
tabla("Defectos detectados y su correccion",
      ["Sintoma observado", "Causa", "Correccion"],
      [["docker compose up falla: bitnami/kafka:3.7 not found",
        "Bitnami retiro sus imagenes versionadas de Docker Hub",
        "Migracion a la imagen oficial apache/kafka:3.9.0 (modo KRaft)"],
       ["Kafka-UI no podria conectarse al broker",
        "advertised.listeners apuntaba a localhost, que dentro del contenedor de Kafka-UI es "
        "el mismo",
        "Dos listeners: EXTERNAL (localhost:9092, para los microservicios) e INTERNAL "
        "(kafka:29092, para Kafka-UI)"],
       ["El productor publicaba pero los consumidores no procesaban y el saldo no cambiaba",
        "El JsonSerializer agrega el header __TypeId__ con la clase del productor; el "
        "consumidor intenta cargarla y falla (not in the trusted packages)",
        "Productor con spring.json.add.type.headers: false y consumidores con "
        "spring.json.use.type.headers: false"],
       ["La fecha viajaba como [2026,9,24]",
        "Jackson serializa LocalDate como arreglo numerico por defecto",
        "@JsonFormat(shape = STRING) en el evento: '2026-09-24'"],
       ["El README afirmaba escalado por consumidores, pero el topico tenia 1 particion",
        "Con una particion un grupo solo puede tener un consumidor activo",
        "Servicio kafka-init en el compose que crea el topico con 3 particiones"]])

h1("9. Instrucciones de ejecucion")
par("Se requieren 6 terminales, mas Docker Compose para el broker:")
code("docker compose up -d                      # Kafka (KRaft), Kafka-UI y creacion del "
     "topico con 3 particiones\n"
     "cd config-server      && mvn spring-boot:run   # 1.\n"
     "cd eureka-server      && mvn spring-boot:run   # 2.\n"
     "cd ms-cuentas         && mvn spring-boot:run   # 3.\n"
     "cd ms-transacciones   && mvn spring-boot:run\n"
     "cd ms-tarjetas        && mvn spring-boot:run\n"
     "cd api-gateway        && mvn spring-boot:run   # 4. al final")
par("Probar el flujo: obtener un token en POST http://localhost:8080/auth/login y registrar "
    "una transaccion con POST http://localhost:8080/transacciones (cuerpo "
    '{"cuentaId":101,"monto":500,"tipo":"credito"}). El detalle completo esta en el README.md.')

h1("10. Conclusion")
bullet("Criterio 1: se definio Publish/Subscribe sobre Kafka, justificado frente a REST "
       "sincrono, colas punto a punto y Event Sourcing, con las decisiones de diseno "
       "(particiones, key, grupos, contrato JSON) y la explicacion del rol del API Gateway "
       "pedida en la retroalimentacion de la Semana 6.")
bullet("Criterio 2: el diagrama muestra el topico, sus particiones, la key, el mensaje con "
       "todos sus campos, los grupos de consumidores y los componentes de soporte.")
bullet("Criterio 3: Resilience4j en los tres microservicios, con el ciclo CLOSED -> OPEN -> "
       "CLOSED probado con una falla real, y resiliencia de la capa de eventos probada con "
       "el broker caido y con un consumidor caido.")
bullet("Criterio 4: el flujo de eventos funciona de punta a punta (saldo 5000 -> 5500 -> "
       "5200) y la escalabilidad se demostro con dos instancias de un grupo repartiendose 12 "
       "eventos sin duplicados.")
par("La verificacion real del sistema permitio encontrar y corregir cinco defectos que la "
    "revision estatica no habria detectado. Quedan documentadas las limitaciones conocidas "
    "(sin Outbox, entrega at-least-once, sin dead-letter) como trabajo futuro.")

h1("11. Referencias")
bullet("Apache Kafka. Apache Kafka Documentation. "
       "https://kafka.apache.org/documentation/")
bullet("Spring. Spring for Apache Kafka Reference. "
       "https://docs.spring.io/spring-kafka/reference/")
bullet("Resilience4j. CircuitBreaker. "
       "https://resilience4j.readme.io/docs/circuitbreaker")
bullet("Spring. Spring Cloud Gateway. "
       "https://docs.spring.io/spring-cloud-gateway/reference/")
bullet("Villagran, K. (s.f.). bank_legacy_data. "
       "https://github.com/KariVillagran/bank_legacy_data")

doc.save(SALIDA)
print("Informe generado en:", SALIDA)
print("Recordar: al abrir en Word, actualizar el Indice (clic derecho -> Actualizar campo -> "
      "Actualizar toda la tabla).")

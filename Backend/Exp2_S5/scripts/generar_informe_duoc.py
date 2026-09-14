# -*- coding: utf-8 -*-
"""
Genera el informe de la Experiencia 2 - Semana 5 (patron BFF, actividad
SUMATIVA individual) sobre la plantilla oficial DUOC
(PBY2202_EFT_Plantilla_Informe_PDF.docx), reutilizando sus estilos y dejando
portada, indice/TOC y aviso legal intactos.

Uso:  python scripts/generar_informe_duoc.py
Requiere la plantilla en el nivel superior de la carpeta S5:
  ../PBY2202_EFT_Plantilla_Informe_PDF.docx
Salida: Informe_Exp2_S5_Felipe_Penaloza.docx (en la raiz del proyecto).
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EVID = os.path.join(RAIZ, "evidencia")
PLANTILLA = os.path.join(os.path.dirname(RAIZ), "PBY2202_EFT_Plantilla_Informe_PDF.docx")
SALIDA = os.path.join(RAIZ, "Informe_Exp2_S5_Felipe_Penaloza.docx")

AUTOR = "Felipe Peñaloza"

if not os.path.exists(PLANTILLA):
    raise SystemExit(f"No se encontro la plantilla en: {PLANTILLA}")

doc = Document(PLANTILLA)
body = doc.element.body

anchor = None
for p in doc.paragraphs:
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
        anchor = p._p
        break
if anchor is None:
    raise SystemExit("No se encontro el parrafo con el salto de seccion (sectPr).")

def set_text(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for r in paragraph.runs[1:]:
        r.text = ""

paras = doc.paragraphs
set_text(paras[16], "Desarrollo Backend III (PBY2203)")
set_text(paras[17], "Implementando el patron arquitectonico Backend for Frontend (BFF)")

p18 = paras[18]
p18.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p18.add_run("BFF Web, Movil y Cajero como servicios independientes para el Banco XYZ")
r.font.size = Pt(14)

p19 = paras[19]
p19.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p19.add_run("Experiencia 2 - Semana 5 (actividad sumativa individual)\n"
                f"{AUTOR}\n"
                "Facilitador disciplinar: Ana Karina Villagran Ibarra")
r.font.size = Pt(11)

INICIO_EJEMPLO, FIN_EJEMPLO = 26, 48
elems_a_borrar = []
idx = 0
for child in list(body):
    tag = child.tag.split("}")[-1]
    if tag in ("p", "tbl"):
        if INICIO_EJEMPLO <= idx <= FIN_EJEMPLO:
            elems_a_borrar.append(child)
        idx += 1
for el in elems_a_borrar:
    el.getparent().remove(el)

def _mover(el):
    anchor.addprevious(el)

def h1(t):
    p = doc.add_paragraph(t, style="Heading 1"); _mover(p._p); return p

def h2(t):
    p = doc.add_paragraph(t, style="Heading 2"); _mover(p._p); return p

def par(t, bold=False):
    p = doc.add_paragraph(); rr = p.add_run(t); rr.bold = bold; _mover(p._p); return p

def bullet(t):
    p = doc.add_paragraph(style="List Paragraph"); p.add_run("•  " + t); _mover(p._p); return p

def code(t):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    rr = p.add_run(t)
    rr.font.name = "Consolas"; rr.font.size = Pt(8.5)
    rr.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    _mover(p._p); return p

def salto_pagina():
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); _mover(p._p)

def imagen(nombre, ancho=6.0, pie=None):
    ruta = os.path.join(EVID, nombre)
    if not os.path.exists(ruta):
        par(f"[FALTA IMAGEN: {nombre}]", bold=True); return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=Inches(ancho)); _mover(p._p)
    if pie:
        cp = doc.add_paragraph(style="titulo-tabla"); cp.add_run(pie); _mover(cp._p)

_tn = [0]
def tabla(titulo, cabeceras, filas):
    _tn[0] += 1
    tp = doc.add_paragraph(style="titulo-tabla")
    tp.add_run(f"Tabla {_tn[0]}: {titulo}"); _mover(tp._p)
    t = doc.add_table(rows=1, cols=len(cabeceras)); t.style = "Table Grid"
    for i, c in enumerate(cabeceras):
        cell = t.rows[0].cells[i]; cell.text = c
        for p in cell.paragraphs:
            for rr in p.runs:
                rr.bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = str(v)
    _mover(t._tbl)
    sp = doc.add_paragraph(); _mover(sp._p)

# ================================================================
#  CONTENIDO
# ================================================================
salto_pagina()

h1("1. Introduccion")
par("Este informe corresponde a la Experiencia 2, Semana 5: la actividad SUMATIVA individual "
    "que continua y profundiza la formativa de la semana anterior sobre el patron Backend for "
    "Frontend (BFF) para el Banco XYZ. La pauta de evaluacion sumativa exige, sobre lo ya "
    "construido, dos cosas que la formativa no exigia con el mismo rigor: que cada BFF sea "
    "realmente independiente (no modulos de un mismo proceso) y que integre y agregue "
    "informacion obtenida desde servicios backend reales, no desde una capa de servicio en el "
    "mismo proceso.")
par("Para cumplir ambos puntos, el proyecto se reestructuro en CUATRO aplicaciones Spring Boot "
    "separadas -banco-servicios, bff-web, bff-movil y bff-cajero-, cada una con su propio "
    "pom.xml, proceso y puerto, comunicadas exclusivamente por HTTP.")

h1("2. Resultado de aprendizaje abordado")
par("RA2. Desarrolla e implementa patrones arquitectonicos conocidos, a partir del diseno, que "
    "den respuesta a diferentes frontends.")
bullet("IL4. Aplica el patron arquitectonico Backend for Frontend (BFF) para la integracion "
       "eficiente entre multiples frontends y backends.")

h1("3. Arquitectura de la solucion: cuatro servicios independientes")
par("La pauta sumativa (criterio 1) distingue explicitamente entre 'implementar BFFs "
    "completamente independientes' (100%) y 'un BFF unico que proporciona respuestas "
    "genericas' (60%) o 'no implementa BFFs especificos, utilizando el backend monolitico "
    "original' (0%). Por eso la solucion no organiza los canales como paquetes de un mismo "
    "proyecto (como en la formativa de la S4), sino como cuatro aplicaciones desplegables por "
    "separado:")
tabla("Los cuatro modulos independientes",
      ["Modulo", "Rol", "Puerto", "Depende de"],
      [["banco-servicios", "Backend real: Cuentas, Transacciones, Tarjetas (BD H2 propia)",
        "9000", "-"],
       ["bff-web", "BFF del canal Web", "8081", "banco-servicios (solo HTTP)"],
       ["bff-movil", "BFF del canal Movil", "8082", "banco-servicios (solo HTTP)"],
       ["bff-cajero", "BFF del canal Cajero (ATM)", "8083", "banco-servicios (solo HTTP)"]])
code("banco-servicios :9000              (backend real, API /api/interno)\n"
     "        ^        ^        ^\n"
     "        | HTTP   | HTTP   | HTTP        (RestClient, sin dependencia de compilacion)\n"
     "        |        |        |\n"
     "  bff-web:8081  bff-movil:8082  bff-cajero:8083\n"
     "  (rol WEB)      (rol MOVIL)     (rol CAJERO)\n"
     "  Cliente Web     Cliente Movil    Cajero ATM")
par("Cada modulo tiene su propio pom.xml con spring-boot-starter-parent como padre directo "
    "(no dependen entre si para compilar), por lo que pueden compilarse, testearse, ejecutarse "
    "y desplegarse de manera totalmente autonoma: 'cd bff-web && mvn spring-boot:run' arranca "
    "ese canal sin necesidad de tocar los otros tres proyectos. Un pom.xml raiz agrupa los "
    "cuatro modulos unicamente como comodidad para compilarlos juntos con un solo comando; no "
    "es un requisito de ejecucion.")
par("Ningun BFF importa una clase de otro BFF ni de banco-servicios: cada uno define su propia "
    "copia de los DTOs de transporte (CuentaDTO, TransaccionDTO, TarjetaDTO) que reflejan el "
    "contrato JSON expuesto por banco-servicios. Es una duplicacion deliberada -sin libreria "
    "compartida- para que ningun BFF quede acoplado en tiempo de compilacion a los demas "
    "modulos: si banco-servicios cambia su codigo interno, ningun BFF necesita recompilarse.")

h1("4. Criterio 1: BFF independiente y optimizado por canal (20 puntos)")
par("Redaccion del nivel Completamente Logrado de la pauta: 'Implementa BFFs completamente "
    "independientes y optimizados para 3 canales, adaptando las respuestas segun las "
    "necesidades especificas de cada uno'.")
bullet("Independencia de PROCESO: bff-web, bff-movil y bff-cajero corren en tres procesos Java "
       "distintos (puertos 8081/8082/8083), cada uno con su propio arranque y su propio log.")
bullet("Independencia de DESPLIEGUE: cada modulo produce su propio JAR ejecutable "
       "(spring-boot-maven-plugin) y su propio pom.xml con spring-boot-starter-parent como "
       "padre; ninguno requiere que otro BFF este compilado para construirse.")
bullet("Independencia de SEGURIDAD: cada BFF define su propia SecurityConfig con un unico "
       "usuario en memoria para su canal. bff-web no conoce las credenciales de movil ni de "
       "cajero (y viceversa): el aislamiento no depende solo de hasRole(), sino de que el "
       "proceso fisicamente no tiene esa informacion.")
bullet("Optimizado por canal: cada uno adapta forma y volumen de la respuesta (ver criterios 2 "
       "y 3).")
par("Verificado en la evidencia (seccion 9): una credencial de otro canal contra un proceso que "
    "no la conoce devuelve 401; y una ruta de otro canal contra un proceso que no la expone "
    "devuelve 404, porque esa ruta directamente no existe en ese proceso.", bold=True)

h1("5. Criterio 2: Optimizacion de respuestas y consumo de recursos (20 puntos)")
par("Redaccion CL: 'Optimiza completamente el tamano de las respuestas y el consumo de "
    "recursos en los 3 canales, alcanzando tiempos de respuesta rapidos'.")
tabla("Tamano de payload por canal para la misma cuenta (101)",
      ["Canal", "Endpoint", "Campos incluidos"],
      [["Web", "GET /api/web/cuentas/101", "id, titular, tipo, saldo, numeroCuenta completo"],
       ["Movil", "GET /api/movil/cuentas/101",
        "id, titular, saldo, cuentaEnmascarada (se omite tipo y numero completo)"],
       ["Cajero", "GET /api/cajero/cuentas/101/saldo",
        "cuentaEnmascarada, saldoDisponible, tarjetasActivas (sin titular ni id)"]])
bullet("Movil ademas limita el historial de transacciones a las 3 mas recientes "
       "(MAX_TX_MOVIL = 3) y omite la descripcion de cada una, reduciendo el volumen frente a "
       "la respuesta completa de Web.")
bullet("Cada BFF hace solo las llamadas HTTP a banco-servicios que su respuesta necesita: el "
       "endpoint /api/cajero/cuentas/{id}/saldo hace 2 llamadas (cuenta + tarjetas), nunca pide "
       "el historial de transacciones que no usa.")
bullet("Consumo de recursos por proceso: al ser aplicaciones separadas, cada canal puede "
       "escalarse o limitarse (memoria, CPU, replicas) de forma independiente segun su carga "
       "real, algo imposible en un monolito con un unico proceso para los tres canales.")

h1("6. Criterio 3: Transformacion y adaptacion de las respuestas (15 puntos)")
par("Redaccion CL: 'Transforma y adapta correctamente las respuestas para los 3 canales, "
    "entregando estructuras y datos especificos de acuerdo con las necesidades de cada uno'.")
tabla("La misma cuenta 101 segun el canal",
      ["Canal", "Endpoint", "Respuesta"],
      [["Web", "GET /api/web/cuentas/101",
        '{"id":101,"titular":"John Doe","tipo":"ahorro","saldo":5000.0,'
        '"numeroCuenta":"1234567890123456"}'],
       ["Movil", "GET /api/movil/cuentas/101",
        '{"id":101,"titular":"John Doe","saldo":5000.0,"cuentaEnmascarada":"**** 3456"}'],
       ["Cajero", "GET /api/cajero/cuentas/101/saldo",
        '{"cuentaEnmascarada":"**** 3456","saldoDisponible":5000.0,"tarjetasActivas":2}']])
bullet("Web (CuentaWebDTO / DashboardWebDTO): estructura completa y anidada, pensada para una "
       "interfaz rica de escritorio.")
bullet("Movil (CuentaMovilDTO / ResumenMovilDTO): estructura plana y reducida, con el numero de "
       "cuenta enmascarado para minimizar datos sensibles en un dispositivo movil.")
bullet("Cajero (SaldoCajeroDTO / RetiroResponse): estructura minima sin datos personales del "
       "titular, apropiada para un espacio publico.")
par("Cada BFF hace esta transformacion mediante un metodo estatico 'desde(...)' en su propio "
    "DTO (por ejemplo CuentaWebDTO.desde(CuentaDTO)), que mapea el DTO de transporte recibido "
    "de banco-servicios a la forma que su canal necesita. El cajero, ademas, calcula un campo "
    "derivado que ningun otro canal expone igual: el conteo de tarjetas activas.")

h1("7. Criterio 4: Integracion y agregacion desde servicios Backend (20 puntos)")
par("Redaccion CL: 'Integra correctamente los BFF con los servicios Backend requeridos, "
    "obteniendo y agregando informacion desde multiples servicios para construir respuestas "
    "funcionales y coherentes para los canales implementados'.")
par("A diferencia de la formativa (donde un CuentaService/TransaccionService/TarjetaService "
    "vivian como @Service inyectados en el mismo proceso del BFF), aqui banco-servicios es un "
    "proceso HTTP aparte y cada BFF lo consume mediante un BancoServiciosClient basado en "
    "RestClient (Spring 6 / Boot 3.4):")
code("public Optional<CuentaDTO> obtenerCuenta(Long id) {\n"
     "    try {\n"
     "        return Optional.ofNullable(restClient.get()\n"
     "            .uri(\"/api/interno/cuentas/{id}\", id)\n"
     "            .retrieve().body(CuentaDTO.class));\n"
     "    } catch (HttpClientErrorException.NotFound e) {\n"
     "        return Optional.empty();\n"
     "    }\n"
     "}")
tabla("Endpoints agregados (combinan mas de una llamada HTTP a banco-servicios)",
      ["BFF", "Endpoint agregado", "Llamadas HTTP que combina"],
      [["Web", "GET /api/web/cuentas/{id}/dashboard",
        "GET /cuentas/{id} + GET /cuentas/{id}/transacciones + GET /cuentas/{id}/tarjetas"],
       ["Movil", "GET /api/movil/cuentas/{id}/resumen",
        "GET /cuentas/{id} + GET /cuentas/{id}/tarjetas + GET /cuentas/{id}/transacciones"],
       ["Cajero", "GET /api/cajero/cuentas/{id}/saldo",
        "GET /cuentas/{id} + GET /cuentas/{id}/tarjetas"]])
par("Prueba de que la agregacion es real (no un valor fijo): la cuenta 103 tiene una unica "
    "tarjeta y esta cargada como INACTIVA en el DataInitializer de banco-servicios "
    "(activa=false). El BFF cajero consulta esa tarjeta por HTTP, la filtra por 'activa' y "
    "devuelve el conteo correcto:", bold=True)
code('GET /api/cajero/cuentas/103/saldo\n'
     '-> {"cuentaEnmascarada":"**** 5678","saldoDisponible":12000.0,"tarjetasActivas":0}')
par("El retiro tambien queda integrado extremo a extremo: la validacion de negocio (monto "
    "positivo, saldo suficiente) vive en banco-servicios; el BFF cajero traduce las respuestas "
    "HTTP 409/400 que recibe a las mismas excepciones que ya manejaba su controlador, "
    "preservando el contrato original con el cliente del cajero.")

h1("8. Criterio 5: Organizacion del codigo para escalabilidad (15 puntos)")
par("Redaccion CL: 'Organiza el codigo de manera modular y estructurado, facil de extender y "
    "escalar sin afectar el codigo existente'.")
code("Exp2_S5_Felipe_Penaloza/\n"
     " |- pom.xml                     <- agregador (solo compila los 4 modulos juntos)\n"
     " |- banco-servicios/            <- backend real, puerto 9000\n"
     " |   |- model / repository / service / config / web (InternoCuentaController)\n"
     " |- bff-web/                    <- BFF Web, puerto 8081\n"
     " |   |- client/     BancoServiciosClient (RestClient) + DTOs de transporte\n"
     " |   |- web/        controller + dto (incluye DashboardWebDTO agregado)\n"
     " |   |- security/   SecurityConfig (solo rol WEB)\n"
     " |- bff-movil/                  <- BFF Movil, puerto 8082  (misma forma)\n"
     " |- bff-cajero/                 <- BFF Cajero, puerto 8083 (incluye /retiro)")
bullet("Agregar un quinto canal (por ejemplo un BFF para un partner externo) significa crear un "
       "quinto modulo Maven con su propio pom.xml y su propio SecurityConfig, sin tocar ninguno "
       "de los cuatro existentes ni recompilarlos.")
bullet("banco-servicios puede evolucionar su modelo de datos o su base de datos sin que los "
       "BFF se enteren, mientras el contrato JSON de /api/interno no cambie: el acoplamiento "
       "entre modulos es unicamente el contrato HTTP, no codigo compartido.")
bullet("Cada BFF puede escalarse horizontalmente (mas instancias) de forma independiente segun "
       "la carga de su canal, sin escalar los demas.")

h1("9. Evidencia de ejecucion")
h2("9.1 Compilacion y tests de los 4 modulos (mvn -f pom.xml)")
comp = os.path.join(EVID, "00_compilacion_test.txt")
if os.path.exists(comp):
    with open(comp, encoding="utf-8", errors="replace") as f:
        code(f.read().strip())
par("Los cuatro modulos compilan y sus tests de carga de contexto Spring pasan "
    "(BUILD SUCCESS, Reactor Summary con los 4 modulos en SUCCESS).")
h2("9.2 Los 4 procesos corriendo simultaneamente (salida real)")
pruebas = os.path.join(EVID, "01_pruebas_bff.txt")
if os.path.exists(pruebas):
    with open(pruebas, encoding="utf-8", errors="replace") as f:
        code(f.read().strip())
par("Se observa, con los 4 procesos levantados a la vez: banco-servicios respondiendo directo "
    "(sin pasar por ningun BFF); el dashboard web agregado; el resumen movil agregado y "
    "enmascarado; el saldo de cajero con tarjetasActivas:0 para la cuenta 103 (agregacion "
    "real); un retiro exitoso y dos rechazados (409 saldo insuficiente, 400 monto invalido); y "
    "las dos formas de aislamiento de canal (401 credencial ajena, 404 ruta que ese proceso no "
    "expone).")

h1("10. Criterio 6: Entrega de los aspectos clave del caso (10 puntos)")
par("Redaccion CL: 'Entrega los 3 aspectos claves solicitados en el caso, para su correcta "
    "funcionalidad'. Los tres aspectos entregados:")
bullet("Codigo fuente: los 4 modulos Maven completos (banco-servicios, bff-web, bff-movil, "
       "bff-cajero), compilables con 'mvn -f pom.xml clean compile'.")
bullet("Documentacion: este informe (arquitectura, justificacion por criterio, instrucciones "
       "de ejecucion) y el README.md del proyecto.")
bullet("Evidencia de ejecucion: capturas reales de compilacion, tests y de los 4 procesos "
       "respondiendo simultaneamente (seccion 9).")

h1("11. Instrucciones de ejecucion")
par("Se requieren 4 terminales, uno por proceso (banco-servicios primero):")
code("cd banco-servicios && mvn spring-boot:run      # arranca en :9000\n"
     "cd bff-web         && mvn spring-boot:run      # arranca en :8081\n"
     "cd bff-movil       && mvn spring-boot:run      # arranca en :8082\n"
     "cd bff-cajero      && mvn spring-boot:run      # arranca en :8083")
par("Compilar y testear los 4 modulos de una vez (sin arrancarlos):")
code("mvn -f pom.xml clean compile\n"
     "mvn -f pom.xml test")
par("Probar la agregacion (PowerShell: usar curl.exe y comillas simples en el body JSON):")
code("curl.exe http://localhost:9000/api/interno/cuentas/103\n"
     "curl.exe -u web_user:web123 http://localhost:8081/api/web/cuentas/101/dashboard\n"
     "curl.exe -u movil_user:movil123 http://localhost:8082/api/movil/cuentas/101/resumen\n"
     "curl.exe -u cajero_user:cajero123 http://localhost:8083/api/cajero/cuentas/103/saldo\n"
     "curl.exe -u cajero_user:cajero123 -X POST http://localhost:8083/api/cajero/cuentas/101/retiro `\n"
     "         -H \"Content-Type: application/json\" -d '{\"monto\": 500}'\n"
     "curl.exe -s -o NUL -w \"%{http_code}\" -u movil_user:movil123 http://localhost:8081/api/web/cuentas  # -> 401")

h1("12. Conclusion")
par("La entrega cubre los seis criterios de la pauta sumativa:")
bullet("Criterio 1 (20 pts): BFFs verdaderamente independientes -cuatro procesos, cuatro "
       "puertos, cuatro despliegues- y optimizados por canal.")
bullet("Criterio 2 (20 pts): tamano de respuesta y llamadas HTTP acotados a lo que cada canal "
       "necesita.")
bullet("Criterio 3 (15 pts): la misma cuenta se transforma en tres estructuras distintas segun "
       "el canal.")
bullet("Criterio 4 (20 pts): cada BFF integra banco-servicios por HTTP y agrega multiples "
       "llamadas en respuestas compuestas, con una prueba verificable (tarjetasActivas:0).")
bullet("Criterio 5 (15 pts): organizacion en modulos Maven independientes, extensible sin "
       "afectar el codigo existente.")
bullet("Criterio 6 (10 pts): codigo fuente, documentacion y evidencia de ejecucion entregados.")

h1("13. Referencias")
bullet("Microsoft. (s.f.). Patron Backends for Frontends. "
       "https://learn.microsoft.com/es-es/azure/architecture/patterns/backends-for-frontends")
bullet("Peplinski, J. (2024). Backend for Frontend (BFF): What You Need to Know. "
       "https://alokai.com/blog/backend-for-frontend")
bullet("Spring. (s.f.). RestClient. "
       "https://docs.spring.io/spring-framework/reference/integration/rest-clients.html")
bullet("Spring. (s.f.). Spring Security Reference - Authorize HttpServletRequests. "
       "https://docs.spring.io/spring-security/reference/servlet/authorization/authorize-http-requests.html")
bullet("Villagran, K. (s.f.). bank_legacy_data. https://github.com/KariVillagran/bank_legacy_data")

doc.save(SALIDA)
print("Informe generado en:", SALIDA)
print("Recordar: al abrir en Word, actualizar el Indice (clic derecho -> Actualizar campo -> "
      "Actualizar toda la tabla).")

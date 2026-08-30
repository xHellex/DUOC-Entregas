# -*- coding: utf-8 -*-
"""
Adapta el contenido del informe a la plantilla oficial DUOC
(PBY2202_EFT_Plantilla_Informe_PDF.docx), reutilizando sus estilos y
manteniendo portada, indice/TOC y aviso legal intactos.

Uso:  python scripts/generar_informe_duoc.py
Requiere la plantilla en el nivel superior de la carpeta S3:
  ../PBY2202_EFT_Plantilla_Informe_PDF.docx
Salida: Informe_Exp1_S3_Felipe_Penaloza.docx (en la raiz del proyecto,
sobrescribe la version anterior sin plantilla).
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EVID = os.path.join(RAIZ, "evidencia")
PLANTILLA = os.path.join(os.path.dirname(RAIZ), "PBY2202_EFT_Plantilla_Informe_PDF.docx")
SALIDA = os.path.join(RAIZ, "Informe_Exp1_S3_Felipe_Penaloza.docx")

if not os.path.exists(PLANTILLA):
    raise SystemExit(f"No se encontro la plantilla en: {PLANTILLA}")

doc = Document(PLANTILLA)
body = doc.element.body

# ---------------------------------------------------------------
# 1) Localizar el parrafo ancla: el ultimo parrafo de la seccion 1
#    (el que trae el sectPr embebido). Todo el contenido nuevo se
#    inserta ANTES de este parrafo, es decir, dentro de la seccion 1,
#    sin tocar portada, indice, TOC ni el aviso legal de la seccion 2.
# ---------------------------------------------------------------
anchor = None
for p in doc.paragraphs:
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
        anchor = p._p
        break
if anchor is None:
    raise SystemExit("No se encontro el parrafo con el salto de seccion (sectPr).")

# ---------------------------------------------------------------
# 2) Completar portada
# ---------------------------------------------------------------
def set_text(paragraph, text):
    """Pone todo el texto en el primer run (conserva su formato) y
    vacia los demas runs, sin tocar el estilo del parrafo."""
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for r in paragraph.runs[1:]:
        r.text = ""

paras = doc.paragraphs
set_text(paras[16], "Desarrollo Backend III (PBY2203)")            # Title
set_text(paras[17], "Optimizando procesos batch para mejorar "
                     "la resiliencia de procesos")                  # Bajada-portada

p18 = paras[18]
p18.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p18.add_run("Escalado y procesamiento paralelo con particionamiento en Spring Batch")
r.font.size = Pt(14)

p19 = paras[19]
p19.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p19.add_run("Experiencia 1 – Semana 3 (actividad sumativa individual)\n"
                "Estudiante: Felipe Penaloza\n"
                "Facilitador disciplinar: Ana Karina Villagran Ibarra")
r.font.size = Pt(11)

# ---------------------------------------------------------------
# 3) Quitar el contenido de ejemplo de la plantilla (indices 26 a 48:
#    Titulo1/Titulo2 de muestra, parrafos lorem ipsum, bullet y las
#    2 tablas de ejemplo), dejando Indice/TOC (20-25) y Aviso legal
#    (seccion 2) intactos.
# ---------------------------------------------------------------
INICIO_EJEMPLO, FIN_EJEMPLO = 26, 48
elems_a_borrar = []
idx = 0
for child in list(body):
    tag = child.tag.split("}")[-1]
    if tag in ("p", "tbl"):
        if INICIO_EJEMPLO <= idx <= FIN_EJEMPLO:
            elems_a_borrar.append(child)
        idx += 1
for el in elems_a_borrar:
    el.getparent().remove(el)

# ---------------------------------------------------------------
# 4) Helpers de contenido: crean el elemento al final del documento
#    (API normal de python-docx) y lo reubican justo antes del ancla,
#    es decir, dentro de la seccion 1, en el orden en que se llaman.
# ---------------------------------------------------------------
def _mover_antes_del_ancla(el):
    anchor.addprevious(el)

def h1(texto):
    p = doc.add_paragraph(texto, style="Heading 1")
    _mover_antes_del_ancla(p._p)
    return p

def h2(texto):
    p = doc.add_paragraph(texto, style="Heading 2")
    _mover_antes_del_ancla(p._p)
    return p

def par(texto, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = bold
    _mover_antes_del_ancla(p._p)
    return p

def bullet(texto):
    p = doc.add_paragraph(style="List Paragraph")
    p.add_run("•  " + texto)
    _mover_antes_del_ancla(p._p)
    return p

def code(texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(texto)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    _mover_antes_del_ancla(p._p)
    return p

def salto_pagina():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    _mover_antes_del_ancla(p._p)

def imagen(nombre, ancho=6.0, pie=None):
    ruta = os.path.join(EVID, nombre)
    if not os.path.exists(ruta):
        par(f"[FALTA IMAGEN: {nombre}]", bold=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=Inches(ancho))
    _mover_antes_del_ancla(p._p)
    if pie:
        cp = doc.add_paragraph(style="titulo-tabla")
        cp.add_run(pie)
        _mover_antes_del_ancla(cp._p)

_tabla_n = [0]
def tabla(titulo, cabeceras, filas):
    _tabla_n[0] += 1
    tp = doc.add_paragraph(style="titulo-tabla")
    tp.add_run(f"Tabla {_tabla_n[0]}: {titulo}")
    _mover_antes_del_ancla(tp._p)

    t = doc.add_table(rows=1, cols=len(cabeceras))
    t.style = "Table Grid"
    for i, c in enumerate(cabeceras):
        cell = t.rows[0].cells[i]
        cell.text = c
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = str(v)
    _mover_antes_del_ancla(t._tbl)
    # espacio despues de la tabla
    sp = doc.add_paragraph()
    _mover_antes_del_ancla(sp._p)

# ================================================================
#  CONTENIDO DEL INFORME
# ================================================================
salto_pagina()

h1("1. Introduccion")
par("Este informe documenta la tercera entrega del proyecto de migracion de los procesos "
    "batch del Banco XYZ. Sobre la base construida en las semanas 1 y 2 (lectura de CSV, "
    "procesamiento por chunks, validaciones, persistencia y procesamiento paralelo con hilos), "
    "la Semana 3 incorpora la tecnica de escalado por particionamiento, la comparacion de "
    "configuraciones para hallar la optima, y refuerzos de resiliencia: limite de omisiones "
    "configurable y jobs idempotentes que se pueden re-ejecutar sin efectos colaterales.")
par("Se recrean los tres procesos legacy como Jobs independientes de Spring Batch:")
bullet("Reporte de Transacciones Diarias: detecta anomalias y genera un resumen agregado.")
bullet("Calculo de Intereses Mensuales: aplica la tasa segun el tipo de cuenta y actualiza el saldo final.")
bullet("Generacion de Estados de Cuenta Anuales: consolida las operaciones por cuenta para auditoria.")

h1("2. Resultado de aprendizaje abordado")
par("RA1. Desarrolla componentes batch para la migracion de sistemas antiguos basados en COBOL/Shell.")
bullet("IL2. Ejecuta jobs y steps en Spring Batch para transformar datos de entrada y generar una salida.")
bullet("IL3. Implementa escalado, procesamiento paralelo y diversas politicas de jobs para optimizar "
       "tiempos de ejecucion, estabilidad y continuidad del sistema.")

h1("3. Arquitectura de la solucion")
par("El proyecto es una aplicacion Spring Boot que orquesta los tres Jobs mediante un "
    "CommandLineRunner (BatchRunner), garantizando el orden de ejecucion y agregando un "
    "parametro timestamp unico por corrida para permitir la reejecucion.")
h2("3.1 Stack tecnologico")
tabla("Stack tecnologico",
      ["Componente", "Tecnologia"],
      [["Framework", "Spring Boot 3.4.1"],
       ["Batch", "Spring Batch 5 (particionamiento, skip/retry, listeners)"],
       ["Lenguaje", "Java 17"],
       ["Persistencia", "Spring Data JPA / Hibernate"],
       ["Base de datos", "MySQL 8 (perfil mysql) / H2 en memoria (perfil h2)"],
       ["Build", "Maven"]])
h2("3.2 Componentes principales")
tabla("Componentes principales",
      ["Clase", "Responsabilidad"],
      [["BatchRunner", "Lanza los 3 Jobs en orden con parametro timestamp unico."],
       ["*BatchConfig (x3)", "Definicion de cada Job: steps, particionamiento y tolerancia a fallos."],
       ["RangoPartitioner", "Divide el total de registros en gridSize rangos (start-end)."],
       ["FlatFileItemReader @StepScope", "Un reader por particion; lee solo su rango del CSV."],
       ["TaskExecutorPartitionHandler", "Ejecuta las particiones en paralelo sobre el TaskExecutor."],
       ["BatchTaskExecutorConfig", "ThreadPoolTaskExecutor; pool = gridSize; prefijo Batch-Part-."],
       ["*Processor (x3)", "Reglas de negocio: validan y transforman cada registro."],
       ["CustomSkipPolicy", "Politica de omision; limite configurable (app.skip.limit)."],
       ["RegistroSkipListener", "Escribe cada omision a salida/errores_<job>.csv."],
       ["JobCompletionListener", "Traza inicio/fin y duracion de cada Job."],
       ["LimpiezaTablasTasklet", "Vacia las tablas destino: hace los Jobs idempotentes."],
       ["Resumen/InformeAnualTasklet", "Steps de consolidacion (Job 1 y Job 3)."],
       ["ContadorRegistros / FechaLegacyParser", "Utilidades: conteo de filas y normalizacion de fechas."]])

h1("4. Procesamiento y validacion de datos")
par("Los datos provienen del repositorio oficial bank_legacy_data y se ubican en "
    "src/main/resources/data/ (conjunto por defecto). Cada Job aplica un ItemProcessor "
    "que valida y corrige los registros; las anomalias lanzan InvalidDataException, lo que "
    "activa la politica de omision y el listener de errores.")
h2("4.1 Reglas de negocio por Job")
par("Job 1 - Transacciones:", bold=True)
bullet("Monto vacio, no numerico, negativo o cero -> anomalia.")
bullet("Tipo distinto de 'debito' o 'credito' -> anomalia.")
bullet("Fecha nula o invalida (ej. 2024-13-01) -> anomalia.")
par("Job 2 - Intereses (tasas: ahorro 0,5% / prestamo 1,5% / hipoteca 0,9%):", bold=True)
bullet("Saldo vacio, no numerico o <= 0 -> anomalia.")
bullet("Edad fuera del rango [18, 75] -> anomalia.")
bullet("Tipo de cuenta desconocido -> anomalia.")
par("Job 3 - Cuentas anuales (prioriza no perder registros de auditoria):", bold=True)
bullet("Fecha invalida -> anomalia (no ubicable en el tiempo).")
bullet("Monto vacio o no numerico -> anomalia.")
bullet("Descripcion faltante -> se completa con 'Sin descripcion' (no se descarta).")
bullet("Monto negativo SI se conserva: representa retiros y compras legitimas.")
h2("4.2 Normalizacion de fechas legacy")
par("FechaLegacyParser acepta los formatos yyyy-MM-dd, yyyy/MM/dd, dd-MM-yyyy y dd/MM/yyyy, "
    "y devuelve null ante valores semanticamente invalidos (mes 13, etc.), que el processor "
    "trata como anomalia.")
h2("4.3 Resultado del procesamiento (datos oficiales)")
tabla("Resultado del procesamiento sobre los datos oficiales",
      ["Job", "Archivo", "Leidas", "Validas", "Anomalias", "Tabla destino (filas)"],
      [["1 Transacciones", "transacciones.csv", 10, 8, 2, "transaccion (8)"],
       ["2 Intereses", "intereses.csv", 8, 6, 2, "cuenta_interes (6)"],
       ["3 Cuentas anuales", "cuentas_anuales.csv", 9, 9, 0, "cuenta_anual (9) / estado_cuenta_anual (8)"]])
par("Anomalias detectadas: transaccion id=3 (monto -200) e id=4 (monto 0); cuenta id=108 "
    "(edad 80) y id=104 (saldo 0). Quedan registradas en salida/errores_transacciones.csv y "
    "salida/errores_intereses.csv.")

h1("5. Manejo de errores y tolerancia a fallos")
par("Cada minionStep (paso worker de una particion) se configura como fault-tolerant y "
    "conserva las tres politicas de resiliencia:")
bullet("skip (CustomSkipPolicy): omite FlatFileParseException e InvalidDataException hasta un "
       "limite configurable con app.skip.limit (por defecto 100.000). El limite es un parametro "
       "mas de la tolerancia a fallos y se ajusta segun el volumen de datos esperado.")
bullet("retry (retryLimit(3), retry(DataAccessException.class)): reintenta fallos transitorios "
       "de acceso a datos antes de omitir.")
bullet("SkipListener (RegistroSkipListener): registra cada omision en el log y en "
       "salida/errores_<job>.csv, indicando fase (lectura/proceso/escritura) y motivo.")
par("Al estar la tolerancia a fallos DENTRO de cada particion, un fallo en una particion se "
    "maneja de forma aislada y no afecta a las demas, lo que aumenta la estabilidad del "
    "sistema batch frente a datos corruptos.")
code("chunk(5) .faultTolerant()\n"
     "        .skipPolicy(customSkipPolicy)\n"
     "        .retryLimit(3).retry(DataAccessException.class)\n"
     "        .listener(new RegistroSkipListener(outputDir, \"transacciones\"))")

h1("6. Politicas de finalizacion, repeticion y reejecucion")
bullet("Finalizacion: JobCompletionListener registra el inicio, el estado final (COMPLETED / "
       "FAILED) y la duracion aproximada de cada Job.")
bullet("Repeticion / reejecucion: BatchRunner agrega un parametro timestamp unico en cada "
       "corrida, de modo que el Job no choca con una instancia previa ya registrada como "
       "COMPLETED en el JobRepository.")
bullet("Idempotencia: cada Job arranca con un step de limpieza (LimpiezaTablasTasklet) que "
       "vacia sus tablas destino con deleteAllInBatch(). Asi, re-ejecutar un Job produce "
       "siempre el mismo resultado, sin acumular filas de corridas anteriores ni distorsionar "
       "los conteos del resumen (total_validas vs total_leidas).")
par("En la evidencia de consola se observa '[LIMPIEZA] 9 fila(s) eliminada(s)...' en la segunda "
    "corrida: el Job detecta y descarta el estado de la corrida anterior antes de procesar.")

h1("7. Escalado por particionamiento (tecnica de la Semana 3)")
par("En la Semana 2 el paralelismo se lograba con un unico step multi-hilo. En la Semana 3 el "
    "step principal de cada Job es un PartitionStep que divide el trabajo en particiones "
    "independientes, cada una ejecutada como un minionStep en su propio hilo.")
code("PartitionStep (paso principal)\n"
     " |- RangoPartitioner: divide N registros en gridSize rangos start-end\n"
     " |- TaskExecutorPartitionHandler: ejecuta las particiones en paralelo\n"
     " |    |- minionStep - particion 0 - hilo Batch-Part-1\n"
     " |    |- minionStep - particion 1 - hilo Batch-Part-2\n"
     " |    |- minionStep - particion 2 - hilo Batch-Part-3\n"
     " |- agrega (aggregate) los resultados")
h2("7.1 Componentes")
bullet("RangoPartitioner: implementa Partitioner. Calcula el tamano de cada particion como "
       "ceil(total / gridSize) y asigna a cada una un ExecutionContext con las claves start y end.")
bullet("Reader @StepScope: cada particion instancia su propio FlatFileItemReader, que lee las "
       "claves start/end de su ExecutionContext y procesa solo ese rango mediante linesToSkip "
       "y maxItemCount. Las particiones no se solapan.")
bullet("TaskExecutorPartitionHandler con setGridSize(app.grid.size): reparte las particiones "
       "sobre el ThreadPoolTaskExecutor (BatchTaskExecutorConfig), cuyo pool se dimensiona igual "
       "a gridSize (un hilo por particion, prefijo de hilo Batch-Part-).")
h2("7.2 Rangos generados (datos oficiales, gridSize = 3)")
tabla("Rangos generados por particion (gridSize = 3)",
      ["Job", "Total registros", "particion0", "particion1", "particion2"],
      [["1 Transacciones", 10, "0-3", "4-7", "8-9"],
       ["2 Intereses", 8, "0-2", "3-5", "6-7"],
       ["3 Cuentas anuales", 9, "0-2", "3-5", "6-8"]])

h1("8. Comparacion de configuraciones y configuracion optima")
par("Para encontrar la configuracion optima se ejecuto el proyecto variando unicamente "
    "app.grid.size (1, 2, 3 y 4) sobre un conjunto de gran volumen generado con "
    "scripts/generar_datos.py: 100.000 filas por archivo, mismo formato y tipos de anomalia "
    "que los datos oficiales. Maquina de 12 nucleos, perfil H2, se reporta el mejor de dos "
    "repeticiones. La duracion de cada Job la entrega el log en la linea "
    "'[JOB] Resumen -> ... duracion aprox: N ms'.")
tabla("Comparacion de tiempos por numero de particiones (100.000 filas/archivo)",
      ["app.grid.size", "Transacciones (ms)", "Intereses (ms)", "Cuentas anuales (ms)", "Total (ms)", "Speedup"],
      [["1 (linea base)", "16.800", "17.119", "16.810", "50.729", "1,00x"],
       ["2", "9.031", "8.684", "8.736", "26.451", "1,92x"],
       ["3", "7.000", "6.692", "6.515", "20.207", "2,51x"],
       ["4", "5.827", "5.231", "5.666", "16.724", "3,03x"]])
par("Eficiencia por particion (speedup / N): 0,96 con 2 particiones; 0,84 con 3; 0,76 con 4.")
h2("8.1 Analisis")
bullet("De 1 a 2 particiones el speedup es casi lineal (1,92x): el particionamiento paga muy "
       "bien porque el trabajo util se reparte sin apenas sobrecosto.")
bullet("De 2 a 3 y de 3 a 4 el tiempo sigue bajando, pero la eficiencia por particion decae "
       "(0,96 -> 0,84 -> 0,76): aparece el costo de coordinar mas hilos y la contencion en la "
       "escritura concurrente sobre la misma base de datos.")
bullet("El 'codo' de la curva esta en gridSize = 3: es el punto donde el retorno por particion "
       "adicional empieza a caer con claridad.")
par("Configuracion optima adoptada: app.grid.size = 3, por ofrecer el mejor equilibrio entre "
    "speedup (2,5x) y eficiencia, sin la complejidad y el desperdicio de recursos de "
    "configuraciones mayores. Es el valor por defecto del proyecto. En una maquina con mas "
    "nucleos y un almacenamiento capaz de absorber mas escritura concurrente, el optimo podria "
    "desplazarse hacia gridSize = 4.")

h1("9. Evidencia de ejecucion")
par("Ejecucion con perfil MySQL sobre los datos oficiales:")
code("mvn clean spring-boot:run \"-Dspring-boot.run.profiles=mysql\"")
h2("9.1 Consola: particionamiento, omisiones y finalizacion")
consola = os.path.join(EVID, "00_consola_ejecucion_mysql.txt")
if os.path.exists(consola):
    with open(consola, encoding="utf-8", errors="replace") as f:
        code(f.read().strip())
par("Se observa: el step [LIMPIEZA] al inicio de cada Job; los rangos [PARTICION] "
    "particion0/1/2; los hilos Batch-Part-1/2/3 ejecutando *MinionStep:particionN en paralelo; "
    "las omisiones [SKIP] '(N de 100000)' con el limite configurable; y los tres Jobs "
    "finalizando en 'Estado: COMPLETED'.")
h2("9.2 Base de datos: el particionamiento registrado")
imagen("01_batch_step_execution_particiones.png", 6.0,
       "Figura 1: tabla batch_step_execution. Cada Job ejecuta limpieza*Step + *Step "
       "(PartitionStep) + tres *MinionStep:particionN, todos COMPLETED.")
imagen("05_conteo_tablas.png", 4.0,
       "Figura 2: conteo de filas por tabla de dominio tras una unica corrida limpia (idempotencia).")
h2("9.3 Datos persistidos")
imagen("02_tabla_transaccion.png", 3.5, "Figura 3: tabla transaccion, 8 transacciones validas (sin acumulacion).")
imagen("03_tabla_resumen_transacciones.png", 6.0,
       "Figura 4: tabla resumen_transacciones (total_leidas=10, total_validas=8, total_anomalias=2).")
imagen("04_tabla_estado_cuenta_anual.png", 5.0,
       "Figura 5: tabla estado_cuenta_anual, consolidado por cuenta.")

h1("10. Instrucciones de ejecucion")
par("Requisito para MySQL: crear la base de datos.")
code("CREATE DATABASE bancoxyz CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;")
par("Ejecucion (los 3 Jobs corren en orden):")
code("# MySQL (PowerShell: comillas obligatorias en el -D)\n"
     "mvn clean spring-boot:run \"-Dspring-boot.run.profiles=mysql\"\n\n"
     "# H2 en memoria\n"
     "mvn clean spring-boot:run")
par("Comparar configuraciones de escalado (editar application.properties):")
code("app.data.path=data/volumen_xl     # 100.000 filas por archivo\n"
     "app.grid.size=1 | 2 | 3 | 4        # numero de particiones\n"
     "app.skip.limit=100000             # limite de omisiones por particion")
par("Regenerar el dataset de volumen:")
code("python scripts/generar_datos.py 100000 src/main/resources/data/volumen_xl")

h1("11. Conclusion")
par("La entrega cumple los requisitos de la actividad sumativa:")
bullet("Los tres procesos legacy estan recreados como Jobs de Spring Batch, leen los CSV "
       "oficiales, aplican validaciones con ItemProcessor y persisten en MySQL.")
bullet("El manejo de errores usa politicas personalizadas (CustomSkipPolicy con limite "
       "configurable) y un SkipListener que deja trazabilidad de cada omision, sin detener el flujo.")
bullet("El escalado se implementa con particionamiento: RangoPartitioner + minionStep @StepScope "
       "+ TaskExecutorPartitionHandler, con el numero de particiones parametrizado.")
bullet("Se compararon cuatro configuraciones de gridSize sobre 100.000 filas y se justifico la "
       "optima (gridSize = 3) a partir de la curva speedup / eficiencia.")
bullet("Las politicas de reejecucion se refuerzan con jobs idempotentes (step de limpieza), de "
       "modo que el sistema se puede volver a ejecutar de forma segura ante un fallo.")
par("El particionamiento reduce el tiempo total de procesamiento de ~50,7 s a ~20,2 s (2,5x) "
    "con la configuracion adoptada, demostrando el beneficio del procesamiento paralelo "
    "distribuido para volumenes grandes, manteniendo la resiliencia del sistema.")

h1("12. Referencias")
bullet("Spring Batch. (s.f.). Scaling and Parallel Processing. "
       "https://docs.spring.io/spring-batch/reference/scalability.html")
bullet("Spring Batch. (s.f.). Spring Batch Architecture. "
       "https://docs.spring.io/spring-batch/reference/spring-batch-architecture.html")
bullet("Villagran, K. (s.f.). bank_legacy_data. https://github.com/KariVillagran/bank_legacy_data")

doc.save(SALIDA)
print("Informe (plantilla DUOC) generado en:", SALIDA)
print("IMPORTANTE: al abrir en Word, actualizar el indice (clic derecho sobre el "
      "Indice -> Actualizar campo -> Actualizar toda la tabla).")

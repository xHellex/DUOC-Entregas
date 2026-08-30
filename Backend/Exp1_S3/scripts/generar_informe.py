# -*- coding: utf-8 -*-
"""
Genera el informe .docx de la Experiencia 1 - Semana 3.
Uso:  python scripts/generar_informe.py
Salida: Informe_Exp1_S3_Felipe_Penaloza.docx (en la raiz del proyecto)
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EVID = os.path.join(RAIZ, "evidencia")
SALIDA = os.path.join(RAIZ, "Informe_Exp1_S3_Felipe_Penaloza.docx")

doc = Document()

# ---- estilos base ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

def h(txt, lvl=1):
    p = doc.add_heading(txt, level=lvl)
    return p

def par(txt, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(txt)
    r.bold = bold
    r.italic = italic
    return p

def bullet(txt):
    doc.add_paragraph(txt, style="List Bullet")

def code(txt):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(txt)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    return p

def imagen(nombre, ancho=6.3, pie=None):
    ruta = os.path.join(EVID, nombre)
    if os.path.exists(ruta):
        doc.add_picture(ruta, width=Inches(ancho))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if pie:
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = c.add_run(pie)
            rr.italic = True
            rr.font.size = Pt(9)
    else:
        par(f"[FALTA IMAGEN: {nombre}]", bold=True)

def tabla(cabeceras, filas):
    t = doc.add_table(rows=1, cols=len(cabeceras))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabeceras):
        cell = t.rows[0].cells[i]
        cell.text = c
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = str(v)
    return t

# ==================== PORTADA ====================
for _ in range(4):
    doc.add_paragraph()
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Optimizando procesos batch para mejorar la resiliencia de procesos")
r.bold = True; r.font.size = Pt(20)

s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Migracion del sistema legacy del Banco XYZ con Spring Batch\nEscalado y procesamiento paralelo mediante particionamiento")
r.font.size = Pt(13)

for _ in range(6):
    doc.add_paragraph()

for linea in [
    "Experiencia 1 - Semana 3 (Actividad sumativa individual)",
    "Desarrollo Backend III (PBY2203)",
    "",
    "Estudiante: Felipe Penaloza",
    "Facilitador disciplinar: Ana Karina Villagran Ibarra",
    "Repositorio: https://github.com/xHellex/DUOC-Entregas/tree/main/Backend/Exp1_S3",
]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(linea).font.size = Pt(11)

doc.add_page_break()

# ==================== 1. INTRODUCCION ====================
h("1. Introduccion", 1)
par("Este informe documenta la tercera entrega del proyecto de migracion de los procesos "
    "batch del Banco XYZ. Sobre la base construida en las semanas 1 y 2 (lectura de CSV, "
    "procesamiento por chunks, validaciones, persistencia y procesamiento paralelo con hilos), "
    "la Semana 3 incorpora la tecnica de escalado por particionamiento, la comparacion de "
    "configuraciones para hallar la optima, y refuerzos de resiliencia: limite de omisiones "
    "configurable y jobs idempotentes que se pueden re-ejecutar sin efectos colaterales.")
par("Se recrean los tres procesos legacy como Jobs independientes de Spring Batch:")
bullet("Reporte de Transacciones Diarias: detecta anomalias y genera un resumen agregado.")
bullet("Calculo de Intereses Mensuales: aplica la tasa segun el tipo de cuenta y actualiza el saldo final.")
bullet("Generacion de Estados de Cuenta Anuales: consolida las operaciones por cuenta para auditoria.")

# ==================== 2. RESULTADO DE APRENDIZAJE ====================
h("2. Resultado de aprendizaje abordado", 1)
par("RA1. Desarrolla componentes batch para la migracion de sistemas antiguos basados en COBOL/Shell.")
bullet("IL2. Ejecuta jobs y steps en Spring Batch para transformar datos de entrada y generar una salida.")
bullet("IL3. Implementa escalado, procesamiento paralelo y diversas politicas de jobs para optimizar "
       "tiempos de ejecucion, estabilidad y continuidad del sistema.")

# ==================== 3. ARQUITECTURA ====================
h("3. Arquitectura de la solucion", 1)
par("El proyecto es una aplicacion Spring Boot que orquesta los tres Jobs mediante un "
    "CommandLineRunner (BatchRunner), garantizando el orden de ejecucion y agregando un "
    "parametro timestamp unico por corrida para permitir la reejecucion.")
h("3.1 Stack tecnologico", 2)
tabla(["Componente", "Tecnologia"],
      [["Framework", "Spring Boot 3.4.1"],
       ["Batch", "Spring Batch 5 (particionamiento, skip/retry, listeners)"],
       ["Lenguaje", "Java 17"],
       ["Persistencia", "Spring Data JPA / Hibernate"],
       ["Base de datos", "MySQL 8 (perfil mysql) / H2 en memoria (perfil h2)"],
       ["Build", "Maven"]])

h("3.2 Componentes principales", 2)
tabla(["Clase", "Responsabilidad"],
      [["BatchRunner", "Lanza los 3 Jobs en orden con parametro timestamp unico."],
       ["*BatchConfig (x3)", "Definicion de cada Job: steps, particionamiento y tolerancia a fallos."],
       ["RangoPartitioner", "Divide el total de registros en gridSize rangos (start-end)."],
       ["FlatFileItemReader @StepScope", "Un reader por particion; lee solo su rango del CSV."],
       ["TaskExecutorPartitionHandler", "Ejecuta las particiones en paralelo sobre el TaskExecutor."],
       ["BatchTaskExecutorConfig", "ThreadPoolTaskExecutor; pool = gridSize; prefijo Batch-Part-."],
       ["*Processor (x3)", "Reglas de negocio: validan y transforman cada registro."],
       ["CustomSkipPolicy", "Politica de omision; limite configurable (app.skip.limit)."],
       ["RegistroSkipListener", "Escribe cada omision a salida/errores_<job>.csv."],
       ["JobCompletionListener", "Traza inicio/fin y duracion de cada Job."],
       ["LimpiezaTablasTasklet", "Vacia las tablas destino: hace los Jobs idempotentes."],
       ["ResumenTransaccionesTasklet / InformeAnualTasklet", "Steps de consolidacion (Job 1 y Job 3)."],
       ["ContadorRegistros / FechaLegacyParser", "Utilidades: conteo de filas y normalizacion de fechas."]])

# ==================== 4. PROCESAMIENTO DE DATOS ====================
h("4. Procesamiento y validacion de datos", 1)
par("Los datos provienen del repositorio oficial bank_legacy_data y se ubican en "
    "src/main/resources/data/ (conjunto por defecto). Cada Job aplica un ItemProcessor "
    "que valida y corrige los registros; las anomalias lanzan InvalidDataException, lo que "
    "activa la politica de omision y el listener de errores.")

h("4.1 Reglas de negocio por Job", 2)
par("Job 1 - Transacciones:", bold=True)
bullet("Monto vacio, no numerico, negativo o cero -> anomalia.")
bullet("Tipo distinto de 'debito' o 'credito' -> anomalia.")
bullet("Fecha nula o invalida (ej. 2024-13-01) -> anomalia.")
par("Job 2 - Intereses (tasas: ahorro 0,5% / prestamo 1,5% / hipoteca 0,9%):", bold=True)
bullet("Saldo vacio, no numerico o <= 0 -> anomalia.")
bullet("Edad fuera del rango [18, 75] -> anomalia.")
bullet("Tipo de cuenta desconocido -> anomalia.")
par("Job 3 - Cuentas anuales (prioriza no perder registros de auditoria):", bold=True)
bullet("Fecha invalida -> anomalia (no ubicable en el tiempo).")
bullet("Monto vacio o no numerico -> anomalia.")
bullet("Descripcion faltante -> se completa con 'Sin descripcion' (no se descarta).")
bullet("Monto negativo SI se conserva: representa retiros y compras legitimas.")

h("4.2 Normalizacion de fechas legacy", 2)
par("FechaLegacyParser acepta los formatos yyyy-MM-dd, yyyy/MM/dd, dd-MM-yyyy y dd/MM/yyyy, "
    "y devuelve null ante valores semanticamente invalidos (mes 13, etc.), que el processor "
    "trata como anomalia.")

h("4.3 Resultado del procesamiento (datos oficiales)", 2)
tabla(["Job", "Archivo", "Leidas", "Validas", "Anomalias", "Tabla destino (filas)"],
      [["1 Transacciones", "transacciones.csv", 10, 8, 2, "transaccion (8)"],
       ["2 Intereses", "intereses.csv", 8, 6, 2, "cuenta_interes (6)"],
       ["3 Cuentas anuales", "cuentas_anuales.csv", 9, 9, 0, "cuenta_anual (9) / estado_cuenta_anual (8)"]])
par("Anomalias detectadas: transaccion id=3 (monto -200) e id=4 (monto 0); cuenta id=108 "
    "(edad 80) y id=104 (saldo 0). Quedan registradas en salida/errores_transacciones.csv y "
    "salida/errores_intereses.csv.")

# ==================== 5. MANEJO DE ERRORES ====================
h("5. Manejo de errores y tolerancia a fallos", 1)
par("Cada minionStep (paso worker de una particion) se configura como fault-tolerant y "
    "conserva las tres politicas de resiliencia:")
bullet("skip (CustomSkipPolicy): omite FlatFileParseException e InvalidDataException hasta un "
       "limite configurable con app.skip.limit (por defecto 100.000). El limite es un parametro "
       "mas de la tolerancia a fallos y se ajusta segun el volumen de datos esperado.")
bullet("retry (retryLimit(3), retry(DataAccessException.class)): reintenta fallos transitorios "
       "de acceso a datos antes de omitir.")
bullet("SkipListener (RegistroSkipListener): registra cada omision en el log y en "
       "salida/errores_<job>.csv, indicando fase (lectura/proceso/escritura) y motivo.")
par("Al estar la tolerancia a fallos DENTRO de cada particion, un fallo en una particion se "
    "maneja de forma aislada y no afecta a las demas, lo que aumenta la estabilidad del "
    "sistema batch frente a datos corruptos.")
code("chunk(5) .faultTolerant()\n"
     "        .skipPolicy(customSkipPolicy)\n"
     "        .retryLimit(3).retry(DataAccessException.class)\n"
     "        .listener(new RegistroSkipListener(outputDir, \"transacciones\"))")

# ==================== 6. POLITICAS DE JOBS ====================
h("6. Politicas de finalizacion, repeticion y reejecucion", 1)
bullet("Finalizacion: JobCompletionListener registra el inicio, el estado final (COMPLETED / "
       "FAILED) y la duracion aproximada de cada Job.")
bullet("Repeticion / reejecucion: BatchRunner agrega un parametro timestamp unico en cada "
       "corrida, de modo que el Job no choca con una instancia previa ya registrada como "
       "COMPLETED en el JobRepository.")
bullet("Idempotencia: cada Job arranca con un step de limpieza (LimpiezaTablasTasklet) que "
       "vacia sus tablas destino con deleteAllInBatch(). Asi, re-ejecutar un Job produce "
       "siempre el mismo resultado, sin acumular filas de corridas anteriores ni distorsionar "
       "los conteos del resumen (total_validas vs total_leidas).")
par("En la evidencia de consola se observa '[LIMPIEZA] 9 fila(s) eliminada(s)...' en la segunda "
    "corrida: el Job detecta y descarta el estado de la corrida anterior antes de procesar.")

# ==================== 7. PARTICIONAMIENTO ====================
h("7. Escalado por particionamiento (tecnica de la Semana 3)", 1)
par("En la Semana 2 el paralelismo se lograba con un unico step multi-hilo. En la Semana 3 el "
    "step principal de cada Job es un PartitionStep que divide el trabajo en particiones "
    "independientes, cada una ejecutada como un minionStep en su propio hilo.")
code("PartitionStep (paso principal)\n"
     " |- RangoPartitioner: divide N registros en gridSize rangos start-end\n"
     " |- TaskExecutorPartitionHandler: ejecuta las particiones en paralelo\n"
     " |    |- minionStep - particion 0 - hilo Batch-Part-1\n"
     " |    |- minionStep - particion 1 - hilo Batch-Part-2\n"
     " |    |- minionStep - particion 2 - hilo Batch-Part-3\n"
     " |- agrega (aggregate) los resultados")
h("7.1 Componentes", 2)
bullet("RangoPartitioner: implementa Partitioner. Calcula el tamano de cada particion como "
       "ceil(total / gridSize) y asigna a cada una un ExecutionContext con las claves start y end.")
bullet("Reader @StepScope: cada particion instancia su propio FlatFileItemReader, que lee las "
       "claves start/end de su ExecutionContext y procesa solo ese rango mediante linesToSkip "
       "y maxItemCount. Las particiones no se solapan.")
bullet("TaskExecutorPartitionHandler con setGridSize(app.grid.size): reparte las particiones "
       "sobre el ThreadPoolTaskExecutor (BatchTaskExecutorConfig), cuyo pool se dimensiona igual "
       "a gridSize (un hilo por particion, prefijo de hilo Batch-Part-).")
h("7.2 Rangos generados (datos oficiales, gridSize = 3)", 2)
tabla(["Job", "Total registros", "particion0", "particion1", "particion2"],
      [["1 Transacciones", 10, "0-3", "4-7", "8-9"],
       ["2 Intereses", 8, "0-2", "3-5", "6-7"],
       ["3 Cuentas anuales", 9, "0-2", "3-5", "6-8"]])

# ==================== 8. COMPARACION DE CONFIGURACIONES ====================
h("8. Comparacion de configuraciones y configuracion optima", 1)
par("Para encontrar la configuracion optima se ejecuto el proyecto variando unicamente "
    "app.grid.size (1, 2, 3 y 4) sobre un conjunto de gran volumen generado con "
    "scripts/generar_datos.py: 100.000 filas por archivo, mismo formato y tipos de anomalia "
    "que los datos oficiales. Maquina de 12 nucleos, perfil H2, se reporta el mejor de dos "
    "repeticiones. La duracion de cada Job la entrega el log en la linea "
    "'[JOB] Resumen -> ... duracion aprox: N ms'.")
tabla(["app.grid.size", "Transacciones (ms)", "Intereses (ms)", "Cuentas anuales (ms)", "Total (ms)", "Speedup"],
      [["1 (linea base)", "16.800", "17.119", "16.810", "50.729", "1,00x"],
       ["2", "9.031", "8.684", "8.736", "26.451", "1,92x"],
       ["3", "7.000", "6.692", "6.515", "20.207", "2,51x"],
       ["4", "5.827", "5.231", "5.666", "16.724", "3,03x"]])
par("Eficiencia por particion (speedup / N): 0,96 con 2 particiones; 0,84 con 3; 0,76 con 4.")
h("8.1 Analisis", 2)
bullet("De 1 a 2 particiones el speedup es casi lineal (1,92x): el particionamiento paga muy "
       "bien porque el trabajo util se reparte sin apenas sobrecosto.")
bullet("De 2 a 3 y de 3 a 4 el tiempo sigue bajando, pero la eficiencia por particion decae "
       "(0,96 -> 0,84 -> 0,76): aparece el costo de coordinar mas hilos y la contencion en la "
       "escritura concurrente sobre la misma base de datos.")
bullet("El 'codo' de la curva esta en gridSize = 3: es el punto donde el retorno por particion "
       "adicional empieza a caer con claridad.")
par("Configuracion optima adoptada: app.grid.size = 3, por ofrecer el mejor equilibrio entre "
    "speedup (2,5x) y eficiencia, sin la complejidad y el desperdicio de recursos de "
    "configuraciones mayores. Es el valor por defecto del proyecto. En una maquina con mas "
    "nucleos y un almacenamiento capaz de absorber mas escritura concurrente, el optimo podria "
    "desplazarse hacia gridSize = 4.")

# ==================== 9. EVIDENCIA ====================
h("9. Evidencia de ejecucion", 1)
par("Ejecucion con perfil MySQL sobre los datos oficiales:")
code("mvn clean spring-boot:run \"-Dspring-boot.run.profiles=mysql\"")

h("9.1 Consola: particionamiento, omisiones y finalizacion", 2)
consola = os.path.join(EVID, "00_consola_ejecucion_mysql.txt")
if os.path.exists(consola):
    with open(consola, encoding="utf-8", errors="replace") as f:
        code(f.read().strip())
par("Se observa: el step [LIMPIEZA] al inicio de cada Job; los rangos [PARTICION] "
    "particion0/1/2; los hilos Batch-Part-1/2/3 ejecutando *MinionStep:particionN en paralelo; "
    "las omisiones [SKIP] '(N de 100000)' con el limite configurable; y los tres Jobs "
    "finalizando en 'Estado: COMPLETED'.")

h("9.2 Base de datos: el particionamiento registrado", 2)
imagen("01_batch_step_execution_particiones.png", 6.0,
       "Tabla batch_step_execution: cada Job ejecuta limpieza*Step + *Step (PartitionStep) + "
       "tres *MinionStep:particionN, todos COMPLETED.")
imagen("05_conteo_tablas.png", 4.2,
       "Conteo de filas por tabla de dominio tras una unica corrida limpia (idempotencia).")

h("9.3 Datos persistidos", 2)
imagen("02_tabla_transaccion.png", 3.6, "Tabla transaccion: 8 transacciones validas (sin acumulacion).")
imagen("03_tabla_resumen_transacciones.png", 6.0,
       "Tabla resumen_transacciones: total_leidas=10, total_validas=8, total_anomalias=2.")
imagen("04_tabla_estado_cuenta_anual.png", 5.2,
       "Tabla estado_cuenta_anual: consolidado por cuenta (movimientos, ingresos, egresos, saldo neto).")

# ==================== 10. EJECUCION ====================
h("10. Instrucciones de ejecucion", 1)
par("Requisito para MySQL: crear la base de datos.")
code("CREATE DATABASE bancoxyz CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;")
par("Ejecucion (los 3 Jobs corren en orden):")
code("# MySQL (PowerShell: comillas obligatorias en el -D)\n"
     "mvn clean spring-boot:run \"-Dspring-boot.run.profiles=mysql\"\n\n"
     "# H2 en memoria\n"
     "mvn clean spring-boot:run")
par("Comparar configuraciones de escalado (editar application.properties):")
code("app.data.path=data/volumen_xl     # 100.000 filas por archivo\n"
     "app.grid.size=1 | 2 | 3 | 4        # numero de particiones\n"
     "app.skip.limit=100000             # limite de omisiones por particion")
par("Regenerar el dataset de volumen:")
code("python scripts/generar_datos.py 100000 src/main/resources/data/volumen_xl")

# ==================== 11. CONCLUSION ====================
h("11. Conclusion", 1)
par("La entrega cumple los requisitos de la actividad sumativa:")
bullet("Los tres procesos legacy estan recreados como Jobs de Spring Batch, leen los CSV "
       "oficiales, aplican validaciones con ItemProcessor y persisten en MySQL.")
bullet("El manejo de errores usa politicas personalizadas (CustomSkipPolicy con limite "
       "configurable) y un SkipListener que deja trazabilidad de cada omision, sin detener el flujo.")
bullet("El escalado se implementa con particionamiento: RangoPartitioner + minionStep @StepScope "
       "+ TaskExecutorPartitionHandler, con el numero de particiones parametrizado.")
bullet("Se compararon cuatro configuraciones de gridSize sobre 100.000 filas y se justifico la "
       "optima (gridSize = 3) a partir de la curva speedup / eficiencia.")
bullet("Las politicas de reejecucion se refuerzan con jobs idempotentes (step de limpieza), de "
       "modo que el sistema se puede volver a ejecutar de forma segura ante un fallo.")
par("El particionamiento reduce el tiempo total de procesamiento de ~50,7 s a ~20,2 s (2,5x) "
    "con la configuracion adoptada, demostrando el beneficio del procesamiento paralelo "
    "distribuido para volumenes grandes, manteniendo la resiliencia del sistema.")

# ==================== 12. REFERENCIAS ====================
h("12. Referencias", 1)
bullet("Spring Batch. (s.f.). Scaling and Parallel Processing. "
       "https://docs.spring.io/spring-batch/reference/scalability.html")
bullet("Spring Batch. (s.f.). Spring Batch Architecture. "
       "https://docs.spring.io/spring-batch/reference/spring-batch-architecture.html")
bullet("Villagran, K. (s.f.). bank_legacy_data. https://github.com/KariVillagran/bank_legacy_data")

doc.save(SALIDA)
print("Informe generado en:", SALIDA)

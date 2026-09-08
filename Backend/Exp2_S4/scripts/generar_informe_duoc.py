# -*- coding: utf-8 -*-
"""
Genera el informe de la Experiencia 2 - Semana 4 (patron BFF) sobre la
plantilla oficial DUOC (PBY2202_EFT_Plantilla_Informe_PDF.docx), reutilizando
sus estilos y dejando portada, indice/TOC y aviso legal intactos.

Uso:  python scripts/generar_informe_duoc.py
Requiere la plantilla en el nivel superior de la carpeta S4:
  ../PBY2202_EFT_Plantilla_Informe_PDF.docx
Salida: Informe_Exp2_S4_Grupo3.docx (en la raiz del proyecto).
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EVID = os.path.join(RAIZ, "evidencia")
PLANTILLA = os.path.join(os.path.dirname(RAIZ), "PBY2202_EFT_Plantilla_Informe_PDF.docx")
SALIDA = os.path.join(RAIZ, "Informe_Exp2_S4_Grupo3.docx")

# Integrantes del grupo:
INTEGRANTES = "Felipe Peñaloza"   # Grupo 3 (integrante unico)

if not os.path.exists(PLANTILLA):
    raise SystemExit(f"No se encontro la plantilla en: {PLANTILLA}")

doc = Document(PLANTILLA)
body = doc.element.body

# --- ancla: ultimo parrafo de la seccion 1 (trae el sectPr embebido) ---
anchor = None
for p in doc.paragraphs:
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
        anchor = p._p
        break
if anchor is None:
    raise SystemExit("No se encontro el parrafo con el salto de seccion (sectPr).")

def set_text(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for r in paragraph.runs[1:]:
        r.text = ""

paras = doc.paragraphs
set_text(paras[16], "Desarrollo Backend III (PBY2203)")
set_text(paras[17], "Aplicando el patron arquitectonico Backend for Frontend (BFF)")

p18 = paras[18]
p18.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p18.add_run("BFF Web, Movil y Cajero para el Banco XYZ")
r.font.size = Pt(14)

p19 = paras[19]
p19.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p19.add_run("Experiencia 2 - Semana 4 (actividad formativa)\n"
                f"Grupo 3 - {INTEGRANTES}\n"
                "Facilitador disciplinar: Ana Karina Villagran Ibarra")
r.font.size = Pt(11)

# --- quitar contenido de ejemplo de la plantilla (indices 26..48) ---
INICIO_EJEMPLO, FIN_EJEMPLO = 26, 48
elems_a_borrar = []
idx = 0
for child in list(body):
    tag = child.tag.split("}")[-1]
    if tag in ("p", "tbl"):
        if INICIO_EJEMPLO <= idx <= FIN_EJEMPLO:
            elems_a_borrar.append(child)
        idx += 1
for el in elems_a_borrar:
    el.getparent().remove(el)

# --- helpers ---
def _mover(el):
    anchor.addprevious(el)

def h1(t):
    p = doc.add_paragraph(t, style="Heading 1"); _mover(p._p); return p

def h2(t):
    p = doc.add_paragraph(t, style="Heading 2"); _mover(p._p); return p

def par(t, bold=False):
    p = doc.add_paragraph(); rr = p.add_run(t); rr.bold = bold; _mover(p._p); return p

def bullet(t):
    p = doc.add_paragraph(style="List Paragraph"); p.add_run("•  " + t); _mover(p._p); return p

def code(t):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    rr = p.add_run(t)
    rr.font.name = "Consolas"; rr.font.size = Pt(8.5)
    rr.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    _mover(p._p); return p

def salto_pagina():
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); _mover(p._p)

def imagen(nombre, ancho=6.0, pie=None):
    ruta = os.path.join(EVID, nombre)
    if not os.path.exists(ruta):
        par(f"[FALTA IMAGEN: {nombre}]", bold=True); return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=Inches(ancho)); _mover(p._p)
    if pie:
        cp = doc.add_paragraph(style="titulo-tabla"); cp.add_run(pie); _mover(cp._p)

_tn = [0]
def tabla(titulo, cabeceras, filas):
    _tn[0] += 1
    tp = doc.add_paragraph(style="titulo-tabla")
    tp.add_run(f"Tabla {_tn[0]}: {titulo}"); _mover(tp._p)
    t = doc.add_table(rows=1, cols=len(cabeceras)); t.style = "Table Grid"
    for i, c in enumerate(cabeceras):
        cell = t.rows[0].cells[i]; cell.text = c
        for p in cell.paragraphs:
            for rr in p.runs:
                rr.bold = True
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = str(v)
    _mover(t._tbl)
    sp = doc.add_paragraph(); _mover(sp._p)

# ================================================================
#  CONTENIDO
# ================================================================
salto_pagina()

h1("1. Introduccion")
par("Este informe corresponde a la Experiencia 2 del proyecto del Banco XYZ. En las semanas "
    "anteriores se trabajo la migracion de procesos legacy con Spring Batch; en esta cuarta "
    "semana se aborda un patron arquitectonico distinto: Backend for Frontend (BFF).")
par("El objetivo es exponer los datos del banco (cuentas, saldos y transacciones) mediante "
    "APIs REST, pero con un backend a medida para cada tipo de cliente: web, movil y cajero "
    "automatico. Cada BFF entrega la misma informacion en el formato y volumen que su cliente "
    "necesita, y con un control de acceso propio por canal.")
par("La actividad pide, ademas de implementar, analizar la estrategia de implementacion mas "
    "conveniente y las ventajas y desventajas del patron.")

h1("2. Resultado de aprendizaje abordado")
par("RA2. Desarrolla e implementa patrones arquitectonicos conocidos, a partir del diseno, que "
    "den respuesta a diferentes frontends.")
bullet("IL4. Aplica el patron arquitectonico Backend for Frontend (BFF) para la integracion "
       "eficiente entre multiples frontends y backends.")

h1("3. El patron Backend for Frontend (BFF)")
par("En una arquitectura tradicional un unico backend generico atiende a todos los frontends "
    "por igual. Esto genera friccion cuando los clientes tienen necesidades distintas: un "
    "navegador de escritorio puede mostrar informacion rica y completa, un movil necesita "
    "respuestas ligeras para ahorrar ancho de banda, y un cajero requiere una interfaz minima "
    "y segura centrada en operaciones criticas. El backend unico termina lleno de condicionales "
    "por tipo de cliente y se vuelve un cuello de botella para los equipos de frontend.")
par("El patron BFF resuelve esto creando un backend especifico para cada frontend. La logica "
    "de negocio se mantiene unica y centralizada; lo que cambia en cada BFF es como se "
    "transforma y expone la respuesta para su cliente.")
h2("3.1 Ventajas")
bullet("Cada frontend recibe exactamente los datos que necesita, en su formato: menos "
       "transformacion en el cliente y menos trafico.")
bullet("Los equipos de frontend pueden evolucionar su BFF sin coordinar con los demas canales.")
bullet("Permite politicas de seguridad y limites distintos por canal (un cajero no necesita el "
       "mismo acceso que la web).")
h2("3.2 Desventajas / cuando NO conviene")
bullet("Mas piezas que mantener y desplegar: si los frontends son pocos y con necesidades "
       "parecidas, un solo backend bien disenado es mas simple.")
bullet("Riesgo de duplicar logica de negocio entre los BFF si no se aisla en una capa comun.")
bullet("Costo de infraestructura si se opta por servicios completamente separados.")
h2("3.3 Escenarios donde resulta mas efectivo")
par("Aplicaciones con varios frontends heterogeneos (web de escritorio + movil + dispositivos "
    "como cajeros o IoT), equipos de frontend independientes, y requisitos de datos o de "
    "seguridad que se contraponen entre canales. Es exactamente el caso del Banco XYZ.")

h1("4. Estrategia de implementacion elegida")
par("El material describe tres estrategias para implementar BFF:")
tabla("Estrategias de implementacion de BFF",
      ["Estrategia", "Descripcion", "Costo / complejidad"],
      [["Backends completamente independientes",
        "Un proyecto y despliegue separado por cada cliente.",
        "Alto: 3 despliegues, 3 configuraciones, riesgo de duplicar logica."],
       ["Endpoints personalizados sobre una base comun",
        "Un unico servicio con un modulo (controller + DTOs) por canal y un nucleo compartido.",
        "Bajo-medio: 1 despliegue, separacion clara por paquetes."],
       ["BFF delgados sobre microservicios",
        "Cada BFF es una capa fina que orquesta varios microservicios de dominio.",
        "Alto: requiere una arquitectura de microservicios previa."]])
par("Estrategia adoptada: endpoints personalizados sobre una base comun (un unico servicio "
    "Spring Boot con modulos separados por canal).", bold=True)
h2("4.1 Justificacion")
bullet("Los tres BFF consumen la MISMA fuente de datos del Banco XYZ; separar en tres "
       "repositorios solo agregaria duplicacion de infraestructura sin beneficio real.")
bullet("Con paquetes separados (bff.web, bff.movil, bff.cajero) sobre un core compartido, la "
       "separacion por cliente queda visible y evaluable, pero se mantiene un solo despliegue "
       "y la logica de negocio no se duplica (vive en core.service.BancoService).")
bullet("No existe una arquitectura de microservicios previa, por lo que la tercera estrategia "
       "no aplica; y para el alcance de la actividad, tres despliegues separados serian "
       "sobreingenieria.")
par("En terminos de patrones de diseno, cada BFF actua como un Adapter: toma el modelo de "
    "dominio de BancoService y lo adapta al DTO que su cliente espera.")

h1("5. Arquitectura de la solucion")
code("                     +------------------------------+\n"
     "  Cliente Web  --->   |  /api/web    (WebBffController)   |--\\\n"
     "                     +------------------------------+   \\\n"
     " Cliente Movil --->   |  /api/movil  (MovilBffController) |----> BancoService ---> BD (H2)\n"
     "                     +------------------------------+   /   (logica unica)\n"
     "   Cajero ATM  --->   |  /api/cajero (CajeroBffController)|--/\n"
     "                     +------------------------------+")
par("Cada BFF tiene su propio controller y sus propios DTOs, y transforma los objetos de "
    "dominio que entrega BancoService (compartido) segun su cliente. BancoService concentra el "
    "acceso a datos (CuentaRepository, TransaccionRepository) y la logica de negocio, incluida "
    "la operacion critica de retiro con validacion de saldo.")
h2("5.1 Stack tecnologico")
tabla("Stack tecnologico",
      ["Componente", "Tecnologia"],
      [["Framework", "Spring Boot 3.4.1"],
       ["APIs REST", "spring-boot-starter-web"],
       ["Seguridad", "spring-boot-starter-security (Basic Auth, usuarios en memoria)"],
       ["Validacion", "spring-boot-starter-validation (jakarta.validation)"],
       ["Persistencia", "Spring Data JPA / Hibernate"],
       ["Base de datos", "H2 en memoria (datos de ejemplo cargados por DataInitializer)"],
       ["Lenguaje / Build", "Java 17 / Maven"]])
h2("5.2 Datos de ejemplo (DataInitializer)")
tabla("Cuentas cargadas al arrancar",
      ["id", "titular", "tipo", "saldo", "numeroCuenta"],
      [["101", "John Doe", "ahorro", "5000", "1234567890123456"],
       ["102", "Jane Smith", "prestamo", "8000", "2345678901234567"],
       ["103", "Bob Johnson", "hipoteca", "12000", "3456789012345678"],
       ["105", "Charlie Green", "ahorro", "7000", "4567890123456789"],
       ["107", "Diana Prince", "prestamo", "15000", "5678901234567890"]])

h1("6. Personalizacion de la informacion por canal")
par("Es el nucleo del patron: la misma cuenta 101, segun el canal que la pida, se entrega en "
    "un formato distinto.")
tabla("La cuenta 101 segun el canal",
      ["Canal", "Endpoint", "Respuesta"],
      [["Web", "GET /api/web/cuentas/101",
        '{"id":101,"titular":"John Doe","tipo":"ahorro","saldo":5000.0,'
        '"numeroCuenta":"1234567890123456"}'],
       ["Movil", "GET /api/movil/cuentas/101",
        '{"id":101,"titular":"John Doe","saldo":5000.0,"cuentaEnmascarada":"**** 3456"}'],
       ["Cajero", "GET /api/cajero/cuentas/101/saldo",
        '{"cuentaEnmascarada":"**** 3456","saldoDisponible":5000.0}']])
bullet("Web (CuentaWebDTO): todos los campos, numero de cuenta completo, historial de "
       "transacciones sin recortar (con id y descripcion). Apto para interfaces complejas.")
bullet("Movil (CuentaMovilDTO): se omite 'tipo', el numero va enmascarado ('**** 3456') y el "
       "historial se limita a las 5 transacciones mas recientes, solo con fecha, monto y tipo. "
       "Menos bytes y menos datos sensibles.")
bullet("Cajero (SaldoCajeroDTO): solo numero enmascarado y saldo disponible; sin nombre del "
       "titular ni otros datos personales, porque el cajero es un espacio publico.")
h2("6.1 Operacion critica del cajero: retiro")
par("El BFF de cajero expone POST /api/cajero/cuentas/{id}/retiro. BancoService.retirar valida "
    "que el monto sea positivo y que haya saldo suficiente antes de descontar:")
tabla("Comportamiento del retiro",
      ["Caso", "Peticion", "Respuesta"],
      [["Retiro valido", 'monto 1000 sobre saldo 5000',
        'HTTP 200 {"exito":true,"mensaje":"Retiro exitoso","saldoRestante":4000.0}'],
       ["Saldo insuficiente", 'monto 999999',
        'HTTP 409 {"exito":false,"mensaje":"Saldo insuficiente","saldoRestante":null}'],
       ["Monto invalido / cuenta inexistente", 'monto <= 0',
        'HTTP 400 {"exito":false,"mensaje":"...","saldoRestante":null}']])

h1("7. Autenticacion y autorizacion por canal")
par("Requisito de la actividad: gestionar autenticacion y autorizacion especificas para cada "
    "canal. Se usa autenticacion basica HTTP con un rol y un usuario por canal (SecurityConfig, "
    "InMemoryUserDetailsManager). Las credenciales de un canal no sirven en otro.")
tabla("Credenciales y acceso por canal",
      ["Canal", "Usuario", "Contrasena", "Rol", "Ruta permitida"],
      [["Web", "web_user", "web123", "WEB", "/api/web/**"],
       ["Movil", "movil_user", "movil123", "MOVIL", "/api/movil/**"],
       ["Cajero", "cajero_user", "cajero123", "CAJERO", "/api/cajero/**"]])
code("http.authorizeHttpRequests(auth -> auth\n"
     "        .requestMatchers(\"/api/web/**\").hasRole(\"WEB\")\n"
     "        .requestMatchers(\"/api/movil/**\").hasRole(\"MOVIL\")\n"
     "        .requestMatchers(\"/api/cajero/**\").hasRole(\"CAJERO\")\n"
     "        .anyRequest().authenticated())\n"
     "     .httpBasic(withDefaults());")
par("Comprobado en la evidencia: credenciales WEB contra /api/movil devuelven 403 Forbidden; "
    "una peticion sin credenciales devuelve 401 Unauthorized; y cada credencial contra su "
    "propia ruta devuelve 200.")

h1("8. Organizacion del codigo segun la estrategia")
par("La estructura de paquetes refleja la estrategia elegida: un nucleo compartido y un modulo "
    "por canal, cada uno con su controller y sus DTOs.")
code("src/main/java/com/bancoxyz/bff/\n"
     " |- BffBancoxyzApplication.java\n"
     " |- core/                     <- nucleo compartido (logica unica)\n"
     " |   |- model/       Cuenta, Transaccion\n"
     " |   |- repository/  CuentaRepository, TransaccionRepository\n"
     " |   |- service/     BancoService     (acceso a datos + retiro)\n"
     " |   |- config/      DataInitializer  (datos de ejemplo)\n"
     " |- web/                      <- BFF Web (respuestas completas)\n"
     " |   |- controller/ WebBffController      (/api/web)\n"
     " |   |- dto/        CuentaWebDTO, TransaccionWebDTO\n"
     " |- movil/                    <- BFF Movil (respuestas ligeras)\n"
     " |   |- controller/ MovilBffController    (/api/movil)\n"
     " |   |- dto/        CuentaMovilDTO, TransaccionMovilDTO\n"
     " |- cajero/                   <- BFF Cajero (operaciones criticas)\n"
     " |   |- controller/ CajeroBffController   (/api/cajero)\n"
     " |   |- dto/        SaldoCajeroDTO, RetiroRequest, RetiroResponse\n"
     " |- security/     SecurityConfig  (un rol por canal)")
bullet("Ningun controller de un canal importa DTOs de otro: la separacion es real.")
bullet("Los tres controllers dependen de la misma clase BancoService: la logica no se duplica.")
bullet("Agregar un cuarto canal (por ejemplo un BFF para un partner externo) es agregar un "
       "paquete mas, sin tocar los existentes.")

h1("9. Evidencia de ejecucion")
h2("9.1 Compilacion y prueba de contexto")
comp = os.path.join(EVID, "00_compilacion_test.txt")
if os.path.exists(comp):
    with open(comp, encoding="utf-8", errors="replace") as f:
        code(f.read().strip())
par("El proyecto compila (18 clases) y el test de carga de contexto Spring pasa "
    "(BUILD SUCCESS, Tests run: 1, Failures: 0).")
h2("9.2 Pruebas de los tres BFF (salida real)")
pruebas = os.path.join(EVID, "01_pruebas_bff.txt")
if os.path.exists(pruebas):
    with open(pruebas, encoding="utf-8", errors="replace") as f:
        code(f.read().strip())
par("Se observa: la misma cuenta 101 en los tres formatos; el retiro valido (saldo 5000 -> "
    "4000) y el rechazado por saldo insuficiente; y los codigos 403 / 401 del aislamiento por "
    "canal.")

h1("10. Instrucciones de ejecucion")
par("Arrancar la aplicacion (H2 en memoria, datos de ejemplo automaticos):")
code("mvn clean spring-boot:run\n"
     "# arranca en http://localhost:8080")
par("Probar los BFF (en PowerShell usar curl.exe; el cuerpo JSON entre comillas simples):")
code("curl.exe -u web_user:web123     http://localhost:8080/api/web/cuentas/101\n"
     "curl.exe -u movil_user:movil123 http://localhost:8080/api/movil/cuentas/101\n"
     "curl.exe -u cajero_user:cajero123 http://localhost:8080/api/cajero/cuentas/101/saldo\n"
     "curl.exe -u cajero_user:cajero123 -X POST http://localhost:8080/api/cajero/cuentas/101/retiro `\n"
     "         -H \"Content-Type: application/json\" -d '{\"monto\": 1000}'\n"
     "curl.exe -s -o NUL -w \"%{http_code}\" -u web_user:web123 http://localhost:8080/api/movil/cuentas/101  # -> 403")

h1("11. Conclusion")
par("La entrega cubre los cuatro criterios de la pauta formativa:")
bullet("Comprension del patron BFF: se explica el problema del backend unico, la solucion, y "
       "las ventajas, desventajas y escenarios de uso (seccion 3).")
bullet("Implementa una estrategia de BFF: se eligio y justifico 'endpoints personalizados sobre "
       "una base comun', frente a las otras dos estrategias (seccion 4).")
bullet("Personaliza la informacion por frontend: la misma cuenta 101 se entrega completa en "
       "web, ligera y enmascarada en movil, y minima en cajero; ademas el cajero suma la "
       "operacion critica de retiro validado (seccion 6), verificado en la evidencia.")
bullet("Organiza el codigo segun la estrategia: nucleo core compartido + un paquete por canal "
       "(web / movil / cajero) + security, sin duplicar logica de negocio (seccion 8).")
par("Adicionalmente se implemento autenticacion y autorizacion por canal (un rol por BFF), "
    "requisito explicito de la actividad, comprobado con los codigos 403 y 401.")

h1("12. Referencias")
bullet("Microsoft. (s.f.). Patron Backends for Frontends. "
       "https://learn.microsoft.com/es-es/azure/architecture/patterns/backends-for-frontends")
bullet("Peplinski, J. (2024). Backend for Frontend (BFF): What You Need to Know. "
       "https://alokai.com/blog/backend-for-frontend")
bullet("Spring. (s.f.). Spring Security Reference - Authorize HttpServletRequests. "
       "https://docs.spring.io/spring-security/reference/servlet/authorization/authorize-http-requests.html")
bullet("Villagran, K. (s.f.). bank_legacy_data. https://github.com/KariVillagran/bank_legacy_data")

doc.save(SALIDA)
print("Informe generado en:", SALIDA)
print("Recordar: al abrir en Word, actualizar el Indice (clic derecho -> Actualizar campo -> "
      "Actualizar toda la tabla) y completar los integrantes del Grupo 3 en la portada.")
